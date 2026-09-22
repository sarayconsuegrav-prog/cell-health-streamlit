"""Interfaz Streamlit para segmentación y conteo de células en fotos y videos.

El video exportado conserva el video original y superpone las máscaras de
segmentación, sin cajas delimitadoras ni etiquetas.
"""

from __future__ import annotations

import csv
from copy import deepcopy
from datetime import datetime
import gc
import hashlib
from html import escape
import json
import math
import os
import re
import shutil
import tempfile
import time
import urllib.request
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from fractions import Fraction
from pathlib import Path
from threading import Lock, Thread
from typing import Any, Callable
from urllib.error import HTTPError
from urllib.parse import quote, urlencode, urlparse

import cv2
import numpy as np
import streamlit as st
from ultralytics import YOLO

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    OPENPYXL_AVAILABLE = True
    OPENPYXL_IMPORT_ERROR = ""
except ImportError as error:
    OPENPYXL_AVAILABLE = False
    OPENPYXL_IMPORT_ERROR = f"{type(error).__name__}: {error}"

try:
    import av
    from streamlit_webrtc import WebRtcMode, webrtc_streamer

    WEBRTC_AVAILABLE = True
    WEBRTC_IMPORT_ERROR = ""
except ImportError as error:
    WEBRTC_AVAILABLE = False
    WEBRTC_IMPORT_ERROR = f"{type(error).__name__}: {error}"


APP_DIR = Path(__file__).resolve().parent
LOCAL_MODEL = APP_DIR / "models" / "best.pt"
MODEL_URL = os.getenv(
    "CELL_MODEL_URL",
    "https://github.com/sarayconsuegrav-prog/cell-health-streamlit/releases/download/v1.0.0/best.pt",
)
MODEL_BACKEND = os.getenv("CELL_MODEL_BACKEND", "openvino").strip().lower()
OPENVINO_MODEL_DIR = APP_DIR / "models" / "best_openvino_model"
OPENVINO_MODEL_URL = os.getenv("CELL_OPENVINO_MODEL_URL", "").strip()
CACHED_MODEL = Path(tempfile.gettempdir()) / "cell-health-streamlit" / "best.pt"
CACHED_OPENVINO_ROOT = Path(tempfile.gettempdir()) / "cell-health-streamlit" / "openvino"
CACHED_OPENVINO_ARCHIVE = CACHED_OPENVINO_ROOT / "best_openvino_model.zip"
PHOTO_RENDER_VERSION = "measurement-diameter-perimeter-class-masks-v11"
# La foto de referencia tiene una barra de escala negra cuyo rótulo indica 50 µm.
SCALE_BAR_LENGTH_UM = 50.0
# Referencia observada en la fotografía entregada con barra: 144 px de barra en 859 px de ancho.
# Se usa únicamente cuando una foto de 40× no conserva la barra visible.
REFERENCE_SCALE_BAR_PX = 144.0
REFERENCE_IMAGE_WIDTH_PX = 859.0
# Las fotos se analizan con más detalle; para cámara/video se usa un tamaño menor
# para que el procesamiento sea ágil en una Mac sin GPU dedicada.
PHOTO_IMAGE_SIZE = 1280
LIVE_IMAGE_SIZE = 640
VIDEO_IMAGE_SIZE = 512
LIVE_INFERENCE_EVERY_N_FRAMES = 3
VIDEO_INFERENCE_STRIDE = 2
VIDEO_OUTPUT_MAX_WIDTH = 1280
LIVE_INFERENCE_RETRY_SECONDS = 2.0
LIVE_RECORD_EVERY_N_FRAMES = 2
LIVE_TRACKER_CONFIG = APP_DIR / "trackers" / "live_botsort.yaml"
LIVE_TRACK_MIN_OBSERVATIONS = 2
# Las mediciones se usan para promedios; conservar cada observación de cada
# fotograma puede hacer crecer la RAM sin mejorar de forma apreciable el reporte.
MAX_MEASUREMENT_OBSERVATIONS = 2000
LIVE_RESOLUTIONS = {
    "640 × 480 (480p)": {
        "width": 640,
        "height": 480,
        "frame_rate": 24,
        "inference_size": 512,
        "inference_every": 3,
    },
    "1280 × 720 (HD)": {
        "width": 1280,
        "height": 720,
        "frame_rate": 20,
        "inference_size": 640,
        "inference_every": 4,
    },
    "1920 × 1080 (Full HD · 2 MP)": {
        "width": 1920,
        "height": 1080,
        "frame_rate": 15,
        "inference_size": 640,
        "inference_every": 5,
    },
}
MAX_LIVE_SAMPLES = 4
POOL_QUERY_KEY = "cell_pools"
NEW_POOL_OPTION = "➕ Registrar nueva piscina"
WHITE_SPOT_LABEL = "Mancha blanca (WSSV)"
VACUOLIZATION_LABEL = "Vacuolización"
HISTORY_LABEL = "Historial"
HISTORY_STATE_KEY = "analysis_history"
HISTORY_VIDEO_ROOT = Path(tempfile.gettempdir()) / "cell-health-streamlit" / "history_videos"
ONEDRIVE_SENT_BATCHES_KEY = "onedrive_sent_batches"
EXCEL_HEADERS = ("Código", "Vacuolización", "White Spot (WSSV)")
# OpenVINO puede no exponer `model.names` en algunas versiones de Ultralytics.
# Estos son los nombres incluidos en `models/best_openvino_model/metadata.yaml`.
DEFAULT_MODEL_NAMES = {
    0: "celula_enferma",
    1: "celula_sana",
}

# Colores BGR de las capas de segmentación.
COLORS_BGR = {
    "sana": (72, 184, 72),       # Verde
    "enferma": (60, 60, 235),    # Rojo
    "otra": (220, 165, 30),      # Azul
}
LIVE_INFERENCE_LOCK = Lock()
DEFAULT_ICE_SERVERS = [
    {"urls": [
        "stun:stun.cloudflare.com:3478",
        "stun:stun.l.google.com:19302",
        "stun:stun1.l.google.com:19302",
        "stun:stun2.l.google.com:19302",
    ]},
]


def runtime_setting(name: str) -> str:
    """Lee una configuración opcional desde Secrets de Streamlit o variables de entorno."""
    environment_value = os.getenv(name, "").strip()
    try:
        secret_value = st.secrets.get(name)
    except Exception:
        secret_value = None
    if secret_value is None:
        return environment_value
    return str(secret_value).strip() or environment_value


class OneDriveConfigurationError(RuntimeError):
    """Indica que falta una configuración necesaria para Microsoft Graph."""


class OneDriveUploadError(RuntimeError):
    """Indica que Microsoft Graph rechazó la lectura o escritura del Excel."""


def _http_error_detail(error: HTTPError) -> str:
    """Obtiene un detalle seguro de Graph sin incluir credenciales."""
    try:
        payload = error.read().decode("utf-8", errors="replace")
    except Exception:
        payload = ""
    if len(payload) > 500:
        payload = payload[:500] + "…"
    return payload or str(error.reason or "sin detalle")


def _get_graph_access_token() -> str:
    """Obtiene un token de aplicación para escribir el archivo del piloto."""
    tenant_id = runtime_setting("MS_TENANT_ID")
    client_id = runtime_setting("MS_CLIENT_ID")
    client_secret = runtime_setting("MS_CLIENT_SECRET")
    if not tenant_id or not client_id or not client_secret:
        raise OneDriveConfigurationError(
            "Configura MS_TENANT_ID, MS_CLIENT_ID y MS_CLIENT_SECRET en "
            "Manage app → Settings → Secrets."
        )

    token_url = f"https://login.microsoftonline.com/{quote(tenant_id, safe='')}/oauth2/v2.0/token"
    payload = urlencode(
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        token_url,
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            token_payload = json.load(response)
    except HTTPError as error:
        raise OneDriveUploadError(
            f"Microsoft no emitió el token ({error.code}): {_http_error_detail(error)}"
        ) from error
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as error:
        raise OneDriveUploadError(f"No se pudo solicitar el token de Microsoft: {error}") from error

    access_token = str(token_payload.get("access_token") or "").strip()
    if not access_token:
        raise OneDriveUploadError("Microsoft devolvió una respuesta sin access_token.")
    return access_token


def _graph_file_content(access_token: str, file_url: str) -> bytes | None:
    """Descarga el Excel desde OneDrive; None significa que todavía no existe."""
    request = urllib.request.Request(
        file_url,
        headers={"Authorization": f"Bearer {access_token}"},
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return response.read()
    except HTTPError as error:
        if error.code == 404:
            return None
        raise OneDriveUploadError(
            f"No se pudo leer el Excel de OneDrive ({error.code}): {_http_error_detail(error)}"
        ) from error
    except OSError as error:
        raise OneDriveUploadError(f"No se pudo conectar con OneDrive: {error}") from error


def _upload_graph_file(access_token: str, file_url: str, workbook_bytes: bytes) -> None:
    """Sube el libro actualizado a la misma ruta de OneDrive."""
    request = urllib.request.Request(
        file_url,
        data=workbook_bytes,
        headers={
            "Authorization": f"Bearer {access_token}",
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        method="PUT",
    )
    try:
        with urllib.request.urlopen(request, timeout=60):
            return
    except HTTPError as error:
        raise OneDriveUploadError(
            f"No se pudo guardar el Excel en OneDrive ({error.code}): {_http_error_detail(error)}"
        ) from error
    except OSError as error:
        raise OneDriveUploadError(f"No se pudo subir el Excel a OneDrive: {error}") from error


def _onedrive_file_url(user: str, file_path: str) -> str:
    """Construye la URL de Graph para un archivo dentro del OneDrive del usuario."""
    normalized_path = "/".join(part for part in file_path.strip("/").split("/") if part)
    if not normalized_path:
        raise OneDriveConfigurationError("MS_ONEDRIVE_FILE_PATH no puede estar vacío.")
    return (
        "https://graph.microsoft.com/v1.0/users/"
        f"{quote(user, safe='')}/drive/root:/{quote(normalized_path, safe='/')}:/content"
    )


@st.cache_data(ttl=300, show_spinner=False)
def fetch_metered_ice_servers(app_name: str, api_key: str) -> list[dict[str, Any]]:
    """Obtiene credenciales TURN temporales de Metered sin exponer la API key."""
    endpoint = (
        f"https://{quote(app_name.strip(), safe='')}.metered.live/"
        f"api/v1/turn/credentials?apiKey={quote(api_key.strip(), safe='')}"
    )
    try:
        with urllib.request.urlopen(endpoint, timeout=8) as response:
            payload = json.load(response)
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return []

    if isinstance(payload, dict):
        payload = payload.get("iceServers", payload.get("ice_servers", []))
    if not isinstance(payload, list):
        return []
    return [server for server in payload if isinstance(server, dict)]


def _filter_cloudflare_ice_servers(payload: Any) -> list[dict[str, Any]]:
    """Normaliza la respuesta de Cloudflare y omite el puerto alternativo 53.

    Cloudflare devuelve varios candidatos TURN para que el navegador elija el
    transporte disponible. El puerto 53 puede quedar bloqueado o tardar en
    expirar en algunos navegadores; no hace falta conservarlo porque la misma
    respuesta incluye 80, 443 y 5349.
    """
    if isinstance(payload, dict):
        payload = payload.get("iceServers", payload.get("ice_servers", []))
    if not isinstance(payload, list):
        return []

    servers: list[dict[str, Any]] = []
    for server in payload:
        if not isinstance(server, dict):
            continue
        urls = server.get("urls")
        if isinstance(urls, str):
            urls = [urls]
        if not isinstance(urls, list):
            continue
        usable_urls = []
        for url in urls:
            if not isinstance(url, str):
                continue
            base_url = url.split("?", 1)[0]
            if base_url.rsplit(":", 1)[-1] == "53":
                continue
            usable_urls.append(url)
        if usable_urls:
            normalized = dict(server)
            normalized["urls"] = usable_urls
            servers.append(normalized)
    return servers


def _normalize_cloudflare_turn_value(value: str) -> str:
    """Limpia valores copiados desde Cloudflare sin alterar el secreto real."""
    normalized = str(value or "").strip()
    if len(normalized) >= 2 and normalized[0] == normalized[-1] and normalized[0] in {'"', "'"}:
        normalized = normalized[1:-1].strip()
    if normalized.lower().startswith("bearer "):
        normalized = normalized[7:].strip()
    return normalized


class CloudflareTurnRequestError(RuntimeError):
    """Error seguro al solicitar credenciales TURN temporales a Cloudflare."""


def _cloudflare_http_error_detail(
    error: HTTPError,
    turn_key_id: str,
    turn_key: str,
) -> str:
    """Extrae solo código/motivo del error y redacta cualquier valor secreto."""
    try:
        raw_body = error.read(8192).decode("utf-8", errors="replace")
    except OSError:
        raw_body = ""
    try:
        payload = json.loads(raw_body) if raw_body else None
    except (json.JSONDecodeError, ValueError):
        payload = None

    errors = payload.get("errors", []) if isinstance(payload, dict) else []
    if isinstance(errors, dict):
        errors = [errors]
    if not isinstance(errors, list):
        errors = []
    if not errors and isinstance(payload, dict):
        for field_name in ("error", "message", "detail", "title"):
            value = payload.get(field_name)
            if isinstance(value, dict):
                errors.append({**value, "code": value.get("code", payload.get("code"))})
            elif isinstance(value, str):
                errors.append({"code": payload.get("code"), "message": value})
    if not errors and raw_body:
        errors = [{"message": raw_body}]

    details = []
    for item in errors[:3]:
        if not isinstance(item, dict):
            continue
        code = item.get("code")
        message = item.get("message")
        safe_parts = []
        if isinstance(code, (int, str)) and str(code).strip():
            safe_parts.append(f"código {str(code)[:24]}")
        if isinstance(message, str) and message.strip():
            safe_message = message.strip()
            for secret in (turn_key, turn_key_id):
                if secret:
                    safe_message = safe_message.replace(secret, "[oculto]")
            safe_message = re.sub(r"(?i)bearer\s+\S+", "Bearer [oculto]", safe_message)
            safe_message = re.sub(r"\b[a-fA-F0-9]{32,}\b", "[valor oculto]", safe_message)
            safe_message = " ".join(safe_message.split())[:180]
            if safe_message:
                safe_parts.append(safe_message)
        if safe_parts:
            details.append(": ".join(safe_parts))
    return " | ".join(details)


@st.cache_data(ttl=300, show_spinner=False)
def _fetch_cloudflare_ice_servers_cached(
    turn_key_id: str,
    turn_key: str,
    ttl_seconds: int,
) -> list[dict[str, Any]]:
    """Solicita credenciales TURN efímeras de Cloudflare.

    La clave larga de TURN solo se usa en el servidor de Streamlit. El
    navegador recibe únicamente el conjunto de iceServers y credenciales con
    caducidad, nunca la clave configurada en Secrets. Solo se cachean respuestas
    exitosas: los errores se elevan como excepciones y Streamlit no los cachea.
    """
    endpoint = (
        "https://rtc.live.cloudflare.com/v1/turn/keys/"
        f"{quote(turn_key_id.strip(), safe='')}/credentials/generate-ice-servers"
    )
    request = urllib.request.Request(
        endpoint,
        data=json.dumps({"ttl": ttl_seconds}).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {turn_key.strip()}",
            "Content-Type": "application/json",
            # El endpoint devuelve 403/1010 al User-Agent predeterminado de
            # urllib; el mismo POST sí se acepta con cURL desde la terminal.
            "User-Agent": "curl/8.7.1",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            payload = json.load(response)
    except HTTPError as error:
        detail = _cloudflare_http_error_detail(error, turn_key_id, turn_key)
        suffix = f" — {detail}" if detail else ""
        if error.code in {401, 403}:
            raise CloudflareTurnRequestError(
                f"Cloudflare rechazó la solicitud TURN (HTTP {error.code}){suffix}."
            ) from error
        raise CloudflareTurnRequestError(
            f"Cloudflare respondió con HTTP {error.code}{suffix}."
        ) from error
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as error:
        raise CloudflareTurnRequestError(
            f"No se pudo consultar Cloudflare ({type(error).__name__})."
        ) from error

    servers = _filter_cloudflare_ice_servers(payload)
    if not servers:
        raise CloudflareTurnRequestError(
            "Cloudflare respondió sin servidores ICE utilizables."
        )
    return servers


def fetch_cloudflare_ice_servers(
    turn_key_id: str,
    turn_key: str,
    ttl_seconds: int,
) -> tuple[list[dict[str, Any]], str]:
    """Devuelve servidores ICE o un error sin cachear las respuestas fallidas."""
    try:
        return (
            _fetch_cloudflare_ice_servers_cached(
                turn_key_id,
                turn_key,
                ttl_seconds,
            ),
            "",
        )
    except CloudflareTurnRequestError as error:
        return [], str(error)


def normalize_metered_app_name(value: str) -> str:
    """Acepta tanto el slug de Metered como su dominio completo."""
    raw_value = value.strip()
    parsed = urlparse(
        raw_value if "://" in raw_value else f"https://{raw_value}"
    )
    host = (parsed.netloc or parsed.path.split("/", 1)[0]).strip().lower()
    suffix = ".metered.live"
    if host.endswith(suffix):
        host = host[: -len(suffix)]
    return host.strip(".")


def rtc_configuration() -> dict[str, Any]:
    """Construye ICE servers con Cloudflare TURN como respaldo opcional.

    La política predeterminada es la de WebRTC: primero se intenta una ruta
    directa (host/srflx) y solo se usa un candidato relay si la red lo exige.
    Metered queda fuera del camino predeterminado para no consumir su cuota;
    puede reactivarse explícitamente con ``RTC_TURN_PROVIDER = "metered"``.
    """
    ice_servers = [dict(server) for server in DEFAULT_ICE_SERVERS]
    turn_setting = runtime_setting("RTC_ENABLE_TURN").lower()
    turn_enabled = turn_setting not in {"0", "false", "no", "off"}
    if not turn_enabled:
        return {"iceServers": ice_servers}

    raw_ice_servers = runtime_setting("RTC_ICE_SERVERS_JSON")
    if raw_ice_servers:
        try:
            parsed = json.loads(raw_ice_servers)
            if isinstance(parsed, dict):
                parsed = parsed.get("iceServers", [])
            if isinstance(parsed, list):
                ice_servers.extend(server for server in parsed if isinstance(server, dict))
        except (TypeError, ValueError, json.JSONDecodeError):
            st.warning("RTC_ICE_SERVERS_JSON no tiene un formato JSON válido; se usará STUN.")
        return {"iceServers": ice_servers}

    turn_provider = runtime_setting("RTC_TURN_PROVIDER").lower() or "cloudflare"
    if turn_provider in {"none", "off", "disabled"}:
        return {"iceServers": ice_servers}

    if turn_provider in {"cloudflare", "cf"}:
        turn_key_id = _normalize_cloudflare_turn_value(
            runtime_setting("CLOUDFLARE_TURN_KEY_ID")
        )
        # Se admite el alias API_TOKEN para facilitar la migración, aunque el
        # valor debe ser la clave larga emitida al crear la TURN key.
        turn_key = _normalize_cloudflare_turn_value(
            runtime_setting("CLOUDFLARE_TURN_KEY") or runtime_setting(
                "CLOUDFLARE_TURN_API_TOKEN"
            )
        )
        ttl_raw = runtime_setting("CLOUDFLARE_TURN_TTL_SECONDS") or "86400"
        try:
            ttl_seconds = max(300, min(int(ttl_raw), 86400))
        except ValueError:
            ttl_seconds = 86400

        if turn_key_id and turn_key:
            if len(turn_key_id) != 32 or len(turn_key) != 64:
                cloudflare_servers, cloudflare_error = [], (
                    "formato inválido: el Key ID debe tener 32 caracteres y "
                    "la clave larga debe tener 64"
                )
            else:
                cloudflare_servers, cloudflare_error = fetch_cloudflare_ice_servers(
                    turn_key_id,
                    turn_key,
                    ttl_seconds,
                )
            if cloudflare_servers:
                ice_servers.extend(cloudflare_servers)
            else:
                detail = cloudflare_error or "respuesta vacía"
                st.warning(
                    f"Cloudflare TURN no está disponible: {detail} "
                    "Se intentará la conexión directa; revisa "
                    "CLOUDFLARE_TURN_KEY_ID y CLOUDFLARE_TURN_KEY."
                )
        elif runtime_setting("METERED_APP_NAME") or runtime_setting("METERED_API_KEY"):
            st.info(
                "Metered está configurado, pero fue desactivado por la migración a "
                "Cloudflare. Configura las credenciales CLOUDFLARE_TURN_* para "
                "habilitar TURN de respaldo."
            )
        return {"iceServers": ice_servers}

    if turn_provider == "metered":
        metered_app_name = normalize_metered_app_name(runtime_setting("METERED_APP_NAME"))
        metered_api_key = runtime_setting("METERED_API_KEY")
        if metered_app_name and metered_api_key:
            if metered_api_key.startswith("pk_live_"):
                st.warning(
                    "La clave de Metered parece ser una clave Publishable de Realtime. "
                    "Usa la API key de TURN REST en METERED_API_KEY."
                )
            else:
                metered_servers = fetch_metered_ice_servers(metered_app_name, metered_api_key)
                if metered_servers:
                    ice_servers.extend(metered_servers)
                else:
                    st.warning(
                        "Metered no devolvió servidores TURN. Verifica METERED_APP_NAME "
                        "y que METERED_API_KEY sea la API key de TURN REST."
                    )
        return {"iceServers": ice_servers}

    if turn_provider in {"custom", "static"}:
        turn_urls = [
            url.strip()
            for url in runtime_setting("RTC_TURN_URLS").split(",")
            if url.strip()
        ]
        turn_username = runtime_setting("RTC_TURN_USERNAME")
        turn_credential = runtime_setting("RTC_TURN_CREDENTIAL")
        if turn_urls and turn_username and turn_credential:
            ice_servers.append(
                {
                    "urls": turn_urls,
                    "username": turn_username,
                    "credential": turn_credential,
                }
            )
        return {"iceServers": ice_servers}

    st.warning(
        f"RTC_TURN_PROVIDER='{turn_provider}' no es válido; se usará conexión directa."
    )
    return {"iceServers": ice_servers}


def is_openvino_model_dir(path: Path) -> bool:
    """Comprueba que una carpeta contiene un modelo OpenVINO completo."""
    return path.is_dir() and any(path.glob("*.xml")) and any(path.glob("*.bin"))


def cached_openvino_model_path() -> Path | None:
    """Obtiene el modelo OpenVINO local o lo descarga desde un ZIP configurado."""
    configured_path = os.getenv("CELL_OPENVINO_MODEL_DIR", "").strip()
    candidates = []
    if configured_path:
        candidates.append(Path(configured_path).expanduser())
    candidates.extend(
        [
            OPENVINO_MODEL_DIR,
            CACHED_OPENVINO_ROOT / "best_openvino_model",
        ]
    )
    for candidate in candidates:
        if is_openvino_model_dir(candidate):
            return candidate

    if not OPENVINO_MODEL_URL:
        return None

    try:
        CACHED_OPENVINO_ROOT.mkdir(parents=True, exist_ok=True)
        if not CACHED_OPENVINO_ARCHIVE.exists():
            with urllib.request.urlopen(OPENVINO_MODEL_URL, timeout=180) as source, CACHED_OPENVINO_ARCHIVE.open("wb") as target:
                shutil.copyfileobj(source, target)

        extract_root = CACHED_OPENVINO_ROOT / "extracted"
        extract_root.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(CACHED_OPENVINO_ARCHIVE) as archive:
            root = extract_root.resolve()
            for member in archive.infolist():
                destination = (extract_root / member.filename).resolve()
                if root not in destination.parents and destination != root:
                    raise ValueError("El ZIP del modelo contiene una ruta no válida.")
            archive.extractall(extract_root)

        for candidate in [extract_root, *extract_root.rglob("*")]:
            if is_openvino_model_dir(candidate):
                return candidate
    except Exception:
        CACHED_OPENVINO_ARCHIVE.unlink(missing_ok=True)
    return None


def default_pt_model_path() -> Path:
    """Usa el modelo local o descarga la copia pública de la release."""
    if LOCAL_MODEL.exists():
        return LOCAL_MODEL
    if CACHED_MODEL.exists():
        return CACHED_MODEL

    try:
        CACHED_MODEL.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(MODEL_URL, timeout=180) as source, CACHED_MODEL.open("wb") as target:
            shutil.copyfileobj(source, target)
        return CACHED_MODEL
    except Exception:
        # La interfaz mostrará el mensaje existente de modelo no encontrado.
        CACHED_MODEL.unlink(missing_ok=True)
        return LOCAL_MODEL


def export_openvino_model(model_path: Path) -> Path | None:
    """Convierte el checkpoint PyTorch a OpenVINO y conserva el resultado en caché."""
    digest = hashlib.sha256()
    with model_path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    model_revision = digest.hexdigest()[:16]
    CACHED_OPENVINO_ROOT.mkdir(parents=True, exist_ok=True)
    cached_source = CACHED_OPENVINO_ROOT / f"{model_path.stem}_{model_revision}.pt"
    exported_dir = CACHED_OPENVINO_ROOT / f"{model_path.stem}_{model_revision}_openvino_model"
    if is_openvino_model_dir(exported_dir):
        return exported_dir

    try:
        if not cached_source.exists():
            shutil.copy2(model_path, cached_source)
        export_model = YOLO(str(cached_source))
        exported_path = export_model.export(
            format="openvino",
            imgsz=LIVE_IMAGE_SIZE,
            dynamic=True,
            nms=False,
            device="cpu",
        )
        exported_path = Path(exported_path)
        if not is_openvino_model_dir(exported_path):
            return None
        if exported_path != exported_dir:
            exported_path.rename(exported_dir)
        return exported_dir
    except Exception:
        return None


def default_model_path() -> Path | None:
    """Usa OpenVINO si fue seleccionado y conserva PyTorch como respaldo."""
    if MODEL_BACKEND in {"openvino", "ov"}:
        if st.session_state.get("openvino_fallback"):
            return default_pt_model_path()
        openvino_path = cached_openvino_model_path()
        if openvino_path is not None:
            return openvino_path
        return export_openvino_model(default_pt_model_path())
    return default_pt_model_path()


def normalize_class_name(name: str) -> str:
    """Agrupa los nombres del modelo en las categorías usadas por el tablero."""
    normalized = name.lower().strip().replace("á", "a").replace("é", "e")
    if "enferm" in normalized:
        return "enferma"
    if "sana" in normalized or "salud" in normalized or "healthy" in normalized:
        return "sana"
    return "otra"


def load_model(model_path: str | Path) -> YOLO:
    """Carga el modelo de YOLO y declara segmentación para carpetas OpenVINO."""
    path = Path(model_path)
    if path.is_dir():
        return YOLO(str(path), task="segment")
    return YOLO(str(path))


def load_live_model(model_path: str | Path) -> YOLO:
    """Carga el modelo de cámara sin depender del estado de Streamlit.

    La cámara se mantiene renderizada mientras este trabajo ocurre en segundo
    plano. Esto evita que el primer clic de «Iniciar detección» bloquee el
    componente WebRTC mientras OpenVINO compila el modelo.
    """
    try:
        loaded_model = load_model(model_path)
    except Exception as error:
        if MODEL_BACKEND in {"openvino", "ov"} and Path(model_path).is_dir():
            fallback_path = default_pt_model_path()
            try:
                loaded_model = load_model(fallback_path)
            except Exception as fallback_error:
                raise RuntimeError(
                    "No fue posible cargar OpenVINO ni el respaldo `.pt`: "
                    f"{fallback_error}"
                ) from error
        else:
            raise

    if loaded_model.task != "segment":
        raise RuntimeError(
            f"El modelo cargado es de tipo `{loaded_model.task}`. "
            "Esta aplicación requiere un modelo de segmentación."
        )
    return loaded_model


def get_model_names(model: YOLO | None) -> dict[int, str]:
    """Obtiene las clases tanto de YOLO/PyTorch como de ciertos backends OpenVINO.

    Algunas versiones de Ultralytics cargan el modelo OpenVINO sin publicar
    `names` a través de `YOLO.__getattr__`. En ese caso usamos las clases
    declaradas en el modelo convertido para que la cámara y los videos sigan
    clasificando las máscaras correctamente.
    """
    if model is None:
        return DEFAULT_MODEL_NAMES.copy()

    candidates: list[Any] = []
    try:
        candidates.append(model.names)
    except Exception:
        # Algunos backends OpenVINO inicializan el predictor al consultar
        # ``names`` y pueden fallar al compilar el modelo en ese momento.
        # La metadata incluida con el modelo ya contiene las clases correctas;
        # no debemos dejar que esa consulta derribe toda la vista en vivo.
        pass

    try:
        backend_model = model.model
    except Exception:
        backend_model = None
    if backend_model is not None:
        try:
            candidates.append(backend_model.names)
        except Exception:
            pass

    for names in candidates:
        if isinstance(names, dict):
            normalized = {}
            for key, value in names.items():
                try:
                    normalized[int(key)] = str(value)
                except (TypeError, ValueError):
                    continue
            if normalized:
                return normalized
        elif isinstance(names, (list, tuple)) and names:
            return {index: str(value) for index, value in enumerate(names)}

    return DEFAULT_MODEL_NAMES.copy()


def get_session_model(model_path: Path) -> YOLO:
    """Conserva el modelo entre reruns sin compartir el tracker entre usuarios."""
    model_key = f"{MODEL_BACKEND}:{model_path}"
    cached_model = st.session_state.get("cell_model")
    cached_model_path = st.session_state.get("cell_model_path")
    if cached_model is None or cached_model_path != model_key:
        cached_model = load_model(model_path)
        st.session_state["cell_model"] = cached_model
        st.session_state["cell_model_path"] = model_key
    return cached_model


def reset_trackers(model: YOLO) -> None:
    """Evita que IDs de un análisis anterior pasen al siguiente video."""
    predictor = getattr(model, "predictor", None)
    for tracker in getattr(predictor, "trackers", []) or []:
        if hasattr(tracker, "reset"):
            tracker.reset()


def _mask_data_at_frame_size(
    result: Any,
    frame_shape: tuple[int, ...],
) -> np.ndarray | None:
    """Devuelve las máscaras rasterizadas, una por instancia, al tamaño del video.

    `Masks.xy` es útil para dibujar contornos, pero en modelos OpenVINO puede
    perder instancias pequeñas al convertir cada máscara a un polígono. La
    matriz rasterizada conserva cada instancia y Ultralytics ya conoce cómo
    retirar el padding de letterbox al escalarla a la imagen original.
    """
    masks = getattr(result, "masks", None)
    raw_data = getattr(masks, "data", None)
    if raw_data is None:
        return None

    height, width = frame_shape[:2]
    try:
        import torch
        from ultralytics.utils import ops

        if isinstance(raw_data, torch.Tensor):
            data = raw_data.detach()
        else:
            data = torch.as_tensor(np.asarray(raw_data))
        if data.ndim == 2:
            data = data.unsqueeze(0)
        if data.ndim != 3:
            return None
        scaled = ops.scale_masks(
            data.unsqueeze(1),
            (height, width),
            mode="nearest",
        ).squeeze(1)
        return scaled.detach().cpu().numpy() > 0.5
    except Exception:
        # El respaldo mantiene funcionando el render si cambia el tipo de
        # tensor entregado por una versión futura del backend.
        try:
            if hasattr(raw_data, "detach"):
                array = raw_data.detach().cpu().numpy()
            else:
                array = np.asarray(raw_data)
            if array.ndim == 2:
                array = array[None, ...]
            if array.ndim == 4 and array.shape[1] == 1:
                array = array[:, 0]
            if array.ndim != 3:
                return None
            return np.stack(
                [
                    cv2.resize(
                        mask.astype(np.float32),
                        (width, height),
                        interpolation=cv2.INTER_NEAREST,
                    )
                    > 0.5
                    for mask in array
                ],
                axis=0,
            )
        except Exception:
            return None


def _paint_instance_mask(
    canvas: np.ndarray,
    mask: np.ndarray,
    color: tuple[int, int, int],
) -> bool:
    """Pinta una instancia y su borde para que células contiguas no se fusionen visualmente."""
    binary_mask = np.asarray(mask, dtype=np.uint8)
    if binary_mask.ndim != 2 or not np.any(binary_mask):
        return False

    canvas[binary_mask.astype(bool)] = color
    contours, _ = cv2.findContours(binary_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if contours:
        outline = tuple(min(255, int(channel) + 45) for channel in color)
        thickness = max(1, round(min(canvas.shape[:2]) / 500))
        cv2.drawContours(canvas, contours, -1, outline, thickness, lineType=cv2.LINE_AA)
    return True


def mask_frame(
    result: Any,
    frame_shape: tuple[int, ...],
    model_names: dict[int, str],
    included_categories: set[str] | None = None,
) -> np.ndarray:
    """Construye un cuadro con una capa visible por cada instancia segmentada."""
    height, width = frame_shape[:2]
    canvas = np.zeros((height, width, 3), dtype=np.uint8)

    if result.masks is None or result.boxes is None:
        return canvas

    class_ids = result.boxes.cls.int().cpu().tolist()
    raster_masks = _mask_data_at_frame_size(result, frame_shape)
    if raster_masks is not None and len(raster_masks) == len(class_ids):
        for mask, class_id in zip(raster_masks, class_ids):
            class_name = model_names.get(int(class_id), str(class_id))
            category = normalize_class_name(class_name)
            if included_categories is not None and category not in included_categories:
                continue
            _paint_instance_mask(canvas, mask, COLORS_BGR[category])
        return canvas

    # Compatibilidad con resultados que solo exponen contornos `.xy`.
    polygons = result.masks.xy
    for polygon, class_id in zip(polygons, class_ids):
        polygon = np.asarray(polygon, dtype=np.int32)
        if polygon.shape[0] < 3:
            continue
        class_name = model_names.get(int(class_id), str(class_id))
        category = normalize_class_name(class_name)
        if included_categories is not None and category not in included_categories:
            continue
        color = COLORS_BGR[category]
        cv2.fillPoly(canvas, [polygon.reshape((-1, 1, 2))], color)
        outline = tuple(min(255, int(channel) + 45) for channel in color)
        cv2.polylines(
            canvas,
            [polygon.reshape((-1, 1, 2))],
            isClosed=True,
            color=outline,
            thickness=max(1, round(min(height, width) / 500)),
            lineType=cv2.LINE_AA,
        )

    return canvas


def overlay_masks(image: np.ndarray, masks: np.ndarray, opacity: float) -> np.ndarray:
    """Superpone las máscaras coloreadas solo donde existe una segmentación."""
    output = image.copy()
    masked_pixels = np.any(masks != 0, axis=2)
    output[masked_pixels] = (
        image[masked_pixels] * (1.0 - opacity) + masks[masked_pixels] * opacity
    ).astype(np.uint8)
    return output


def draw_live_status_overlay(
    image: np.ndarray,
    detection_active: bool,
    copy_output: bool = True,
) -> np.ndarray:
    """Dibuja una señal visible sobre el video para indicar el estado actual."""
    output = image.copy() if copy_output else image
    height, width = output.shape[:2]
    if not detection_active:
        return output
    label = "EN VIVO"
    badge_color = (40, 40, 220)  # rojo en BGR

    font = cv2.FONT_HERSHEY_SIMPLEX
    # Escalar junto con la resolución de entrada para que el rótulo conserve
    # prácticamente el mismo tamaño entre 480p, HD y Full HD.
    display_scale = max(width / 640.0, 0.75)
    font_scale = 0.32 * display_scale
    thickness = max(1, round(font_scale * 2))
    (text_width, text_height), baseline = cv2.getTextSize(
        label, font, font_scale, thickness
    )
    left = max(8, round(8 * display_scale))
    top = max(8, round(8 * display_scale))
    padding_x = max(5, round(7 * display_scale))
    padding_y = max(4, round(5 * display_scale))
    badge_height = text_height + baseline + padding_y * 2
    badge_width = text_width + padding_x * 3 + round(badge_height * 0.35)
    right = min(width - 8, left + badge_width)
    bottom = min(height - 8, top + badge_height)

    background = output.copy()
    cv2.rectangle(background, (left, top), (right, bottom), (5, 24, 47), -1)
    output = cv2.addWeighted(background, 0.82, output, 0.18, 0)
    center = (left + padding_x + round(badge_height * 0.17), top + badge_height // 2)
    radius = max(5, round(badge_height * 0.16))
    cv2.circle(output, center, radius, badge_color, -1, lineType=cv2.LINE_AA)
    cv2.putText(
        output,
        label,
        (left + padding_x * 2 + round(badge_height * 0.24), top + padding_y + text_height),
        font,
        font_scale,
        (255, 255, 255),
        thickness,
        lineType=cv2.LINE_AA,
    )
    cv2.rectangle(output, (left, top), (right, bottom), badge_color, 1, lineType=cv2.LINE_AA)
    return output


def update_track_votes(result: Any, votes: dict[int, Counter], model_names: dict[int, str]) -> int:
    """Registra la clase observada para cada ID único producido por el tracker."""
    if result.boxes is None or result.boxes.id is None:
        return 0

    track_ids = result.boxes.id.int().cpu().tolist()
    class_ids = result.boxes.cls.int().cpu().tolist()
    for track_id, class_id in zip(track_ids, class_ids):
        # Ultralytics puede emitir -1 para detecciones sin identidad asignada;
        # no deben fusionarse entre sí ni aumentar el conteo acumulado.
        if int(track_id) < 0:
            continue
        class_name = model_names.get(int(class_id), str(class_id))
        votes[int(track_id)][normalize_class_name(class_name)] += 1
    return sum(int(track_id) >= 0 for track_id in track_ids)


def summarize_tracks(
    votes: dict[int, Counter],
    min_observations: int = 1,
) -> dict[str, int]:
    """Cuenta IDs confirmados según la clase observada en más fotogramas."""
    counts = Counter()
    for class_votes in votes.values():
        if sum(class_votes.values()) < min_observations:
            continue
        category, _ = class_votes.most_common(1)[0]
        counts[category] += 1
    return {
        "sana": counts["sana"],
        "enferma": counts["enferma"],
        "otra": counts["otra"],
        "total": sum(counts.values()),
    }


def summarize_instances(result: Any, model_names: dict[int, str]) -> dict[str, int]:
    """Cuenta las instancias detectadas en una sola fotografía."""
    counts = Counter()
    if result.boxes is not None:
        for class_id in result.boxes.cls.int().cpu().tolist():
            class_name = model_names.get(int(class_id), str(class_id))
            counts[normalize_class_name(class_name)] += 1
    return {
        "sana": counts["sana"],
        "enferma": counts["enferma"],
        "otra": counts["otra"],
        "total": sum(counts.values()),
    }


def detect_scale_bar_length_px(image: np.ndarray) -> float | None:
    """Encuentra la barra negra horizontal de escala en la parte superior izquierda."""
    height, width = image.shape[:2]
    if height < 40 or width < 100:
        return None

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    # La barra de referencia proporcionada está arriba a la izquierda. La apertura
    # horizontal elimina el texto y conserva únicamente segmentos largos y rectos.
    search_height = max(1, int(height * 0.30))
    search_width = max(1, int(width * 0.40))
    search_area = gray[:search_height, :search_width]
    dark_pixels = cv2.threshold(search_area, 100, 255, cv2.THRESH_BINARY_INV)[1]
    kernel_width = max(20, int(round(width * 0.02)))
    horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (kernel_width, 1))
    horizontal_lines = cv2.morphologyEx(dark_pixels, cv2.MORPH_OPEN, horizontal_kernel)
    contours, _ = cv2.findContours(horizontal_lines, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    minimum_length = max(40, int(round(width * 0.04)))
    candidates: list[tuple[float, float]] = []
    for contour in contours:
        x, y, line_width, line_height = cv2.boundingRect(contour)
        if line_width < minimum_length or line_height > 8:
            continue
        aspect_ratio = line_width / max(line_height, 1)
        if aspect_ratio < 12:
            continue
        candidates.append((float(line_width), aspect_ratio))

    if not candidates:
        return None
    # La barra es el segmento horizontal más largo que cumple las restricciones.
    return max(candidates, key=lambda candidate: (candidate[0], candidate[1]))[0]


def calibration_from_image(image: np.ndarray) -> dict[str, Any]:
    """Obtiene la calibración directa o estimada de una imagen a 40×."""
    scale_bar_length_px = detect_scale_bar_length_px(image)
    if scale_bar_length_px is None:
        # Algunas fotos del conjunto fueron exportadas sin la barra visible. Al
        # conservar 40× y el mismo campo de visión, la escala se ajusta a su ancho.
        scale_bar_length_px = image.shape[1] * REFERENCE_SCALE_BAR_PX / REFERENCE_IMAGE_WIDTH_PX
        calibration_method = "estimada a 40×"
    else:
        calibration_method = "barra de escala detectada"
    micrometers_per_pixel = SCALE_BAR_LENGTH_UM / scale_bar_length_px
    return {
        "micrometers_per_pixel": micrometers_per_pixel,
        "scale_bar_length_px": scale_bar_length_px,
        "calibration_method": calibration_method,
    }


def new_measurement_samples() -> dict[str, dict[str, list[float]]]:
    """Prepara las listas de métricas de las máscaras segmentadas."""
    return {
        "areas_px2": {"sana": [], "enferma": []},
        "perimeters_px": {"sana": [], "enferma": []},
        "diameters_px": {"sana": [], "enferma": []},
    }


@dataclass
class LiveSessionState:
    """Estado compartido entre el callback WebRTC y la interfaz Streamlit."""

    lock: Lock = field(default_factory=Lock)
    detection_active: bool = False
    frame_number: int = 0
    camera_frames: int = 0
    captured_frames: int = 0
    processed_frames: int = 0
    recording_writer: Any | None = None
    recording_path: Path | None = None
    recording_size: tuple[int, int] | None = None
    recorded_frames: int = 0
    track_votes: dict[int, Counter] = field(default_factory=lambda: defaultdict(Counter))
    measurement_track_ids: set[int] = field(default_factory=set)
    measurement_samples: dict[str, dict[str, list[float]]] = field(default_factory=new_measurement_samples)
    calibration: dict[str, Any] | None = None
    cached_shape: tuple[int, ...] | None = None
    cached_healthy_masks: np.ndarray | None = None
    cached_sick_masks: np.ndarray | None = None
    inference_busy: bool = False
    inference_generation: int = 0
    last_error: str = ""
    inference_retry_at: float = 0.0
    model_loading: bool = False
    inference_model: YOLO | None = None
    inference_model_names: dict[int, str] = field(default_factory=lambda: DEFAULT_MODEL_NAMES.copy())
    completed_samples: list[dict[str, Any]] = field(default_factory=list)
    pending_sample: dict[str, Any] | None = None
    sample_ready: bool = False

    def _close_recording_locked(self, delete_file: bool = False) -> None:
        """Cierra el archivo temporal actual sin borrar un video ya detenido."""
        if self.recording_writer is not None:
            self.recording_writer.release()
            self.recording_writer = None
        if delete_file and self.recording_path is not None:
            try:
                self.recording_path.unlink(missing_ok=True)
            except OSError:
                pass
            self.recording_path = None
        self.recording_size = None

    def record_frame(self, frame: np.ndarray, fps: float) -> None:
        """Escribe el fotograma mostrado en un MP4 temporal de baja carga."""
        with self.lock:
            if self.recording_writer is None:
                recording_file = tempfile.NamedTemporaryFile(suffix=".mp4", delete=False)
                self.recording_path = Path(recording_file.name)
                recording_file.close()
                output_width = min(int(frame.shape[1]), VIDEO_OUTPUT_MAX_WIDTH)
                output_height = max(1, round(frame.shape[0] * output_width / frame.shape[1]))
                self.recording_size = (output_width, output_height)
                self.recording_writer = cv2.VideoWriter(
                    str(self.recording_path),
                    cv2.VideoWriter_fourcc(*"mp4v"),
                    max(float(fps), 1.0),
                    self.recording_size,
                )
                if not self.recording_writer.isOpened():
                    self._close_recording_locked(delete_file=True)
                    self.last_error = "No se pudo crear el video temporal de la muestra."
                    return

            output_frame = frame
            if output_frame.shape[1] != self.recording_size[0] or output_frame.shape[0] != self.recording_size[1]:
                output_frame = cv2.resize(
                    output_frame,
                    self.recording_size,
                    interpolation=cv2.INTER_AREA,
                )
            try:
                self.recording_writer.write(output_frame)
                self.recorded_frames += 1
            except Exception as error:
                self.last_error = f"No se pudo grabar un fotograma: {error}"

    def reset_metrics(self) -> None:
        """Limpia el conteo al iniciar una nueva captura."""
        with self.lock:
            self._close_recording_locked(delete_file=True)
            self.frame_number = 0
            self.camera_frames = 0
            self.captured_frames = 0
            self.processed_frames = 0
            self.recorded_frames = 0
            self.track_votes = defaultdict(Counter)
            self.measurement_track_ids = set()
            self.measurement_samples = new_measurement_samples()
            self.calibration = None
            self.cached_shape = None
            self.cached_healthy_masks = None
            self.cached_sick_masks = None
            self.inference_generation += 1
            self.inference_busy = False
            self.last_error = ""
            self.inference_retry_at = 0.0
            self.pending_sample = None
            self.sample_ready = False

    def discard_pending_sample(self) -> None:
        """Descarta la captura detenida y elimina su video temporal."""
        with self.lock:
            pending_path = ""
            if self.pending_sample:
                pending_path = str(self.pending_sample.get("recording_path") or "").strip()

        # reset_metrics elimina el archivo que todavía está asociado al estado;
        # el segundo borrado cubre el caso en que Streamlit haya reconstruido el
        # estado desde `live_pending_sample` después de un rerun.
        self.reset_metrics()
        if pending_path:
            try:
                Path(pending_path).unlink(missing_ok=True)
            except OSError:
                pass

    def _freeze_current_sample_locked(self) -> None:
        """Congela el resultado al detener para que el guardado no dependa del callback."""
        if self.captured_frames <= 0:
            self.pending_sample = None
            self.sample_ready = False
            return

        measurements = (
            summarize_measurement_samples(self.measurement_samples, self.calibration)
            if self.calibration is not None
            else None
        )
        self.pending_sample = {
            "captured_frames": self.captured_frames,
            "processed_frames": self.processed_frames,
            "recorded_frames": self.recorded_frames,
            "recording_path": str(self.recording_path) if self.recording_path else "",
            "counts": dict(
                summarize_tracks(
                    self.track_votes,
                    min_observations=LIVE_TRACK_MIN_OBSERVATIONS,
                )
            ),
            "measurements": dict(measurements) if measurements is not None else None,
        }
        self.sample_ready = True

    def set_detection_active(self, active: bool) -> None:
        with self.lock:
            self.detection_active = active
            if not active:
                self.inference_busy = False
                self._close_recording_locked()
                self._freeze_current_sample_locked()

    def clear_completed_samples(self) -> None:
        with self.lock:
            history_paths = {
                str(record.get("recording_path") or "").strip()
                for record in get_analysis_history()
                if record.get("recording_path")
            }
            for sample in self.completed_samples:
                recording_path = sample.get("recording_path")
                if recording_path and str(recording_path) not in history_paths:
                    try:
                        Path(recording_path).unlink(missing_ok=True)
                    except OSError:
                        pass
            self.completed_samples = []

    def save_current_sample(self, sample_code: str, lot_name: str) -> bool:
        """Guarda las métricas acumuladas como una de las cuatro muestras."""
        with self.lock:
            if len(self.completed_samples) >= MAX_LIVE_SAMPLES:
                return False

            if self.pending_sample is None:
                self._freeze_current_sample_locked()
            if self.pending_sample is None:
                return False

            pending_sample = self.pending_sample
            sample_number = len(self.completed_samples) + 1
            self.completed_samples.append(
                {
                    "sample_number": sample_number,
                    "label": f"Camarón muestra {sample_number}",
                    "code": sample_code.strip() or f"Muestra-{sample_number}",
                    "lot_name": lot_name.strip() or "Lote sin nombre",
                    "captured_frames": pending_sample["captured_frames"],
                    "processed_frames": pending_sample["processed_frames"],
                    "recorded_frames": pending_sample["recorded_frames"],
                    "recording_path": pending_sample["recording_path"],
                    "counts": dict(pending_sample["counts"]),
                    "measurements": (
                        dict(pending_sample["measurements"])
                        if pending_sample.get("measurements") is not None
                        else None
                    ),
                }
            )
            # El archivo pasa a ser propiedad de la muestra guardada. El
            # siguiente análisis solo podrá borrar su propio archivo temporal.
            self.recording_path = None
            self.recording_size = None
            self.frame_number = 0
            self.captured_frames = 0
            self.processed_frames = 0
            self.recorded_frames = 0
            self.track_votes = defaultdict(Counter)
            self.measurement_track_ids = set()
            self.measurement_samples = new_measurement_samples()
            self.calibration = None
            self.cached_shape = None
            self.cached_healthy_masks = None
            self.cached_sick_masks = None
            self.inference_generation += 1
            self.inference_busy = False
            self.last_error = ""
            self.inference_retry_at = 0.0
            self.pending_sample = None
            self.sample_ready = False
            return True

    def completed_samples_snapshot(self) -> list[dict[str, Any]]:
        """Devuelve una copia de las muestras finalizadas para renderizar y descargar."""
        with self.lock:
            return [
                {
                    **sample,
                    "counts": dict(sample["counts"]),
                    "measurements": (
                        dict(sample["measurements"])
                        if sample.get("measurements") is not None
                        else None
                    ),
                }
                for sample in self.completed_samples
            ]

    def update_completed_sample_code(self, sample_number: int, code: str) -> bool:
        """Corrige el código de una muestra ya guardada."""
        normalized_code = code.strip() or f"Muestra-{sample_number}"
        with self.lock:
            for sample in self.completed_samples:
                if sample.get("sample_number") == sample_number:
                    sample["code"] = normalized_code
                    return True
        return False

    def snapshot(self) -> dict[str, Any]:
        """Obtiene una copia coherente de las métricas para el panel lateral."""
        with self.lock:
            pending_sample = self.pending_sample if self.sample_ready else None
            if pending_sample is not None:
                counts = dict(pending_sample["counts"])
                measurements = (
                    dict(pending_sample["measurements"])
                    if pending_sample.get("measurements") is not None
                    else None
                )
                captured_frames = int(pending_sample["captured_frames"])
                processed_frames = int(pending_sample["processed_frames"])
                recorded_frames = int(pending_sample["recorded_frames"])
            else:
                counts = summarize_tracks(
                    self.track_votes,
                    min_observations=LIVE_TRACK_MIN_OBSERVATIONS,
                )
                measurements = (
                    summarize_measurement_samples(self.measurement_samples, self.calibration)
                    if self.calibration is not None
                    else None
                )
                captured_frames = self.captured_frames
                processed_frames = self.processed_frames
                recorded_frames = self.recorded_frames
            return {
                "detection_active": self.detection_active,
                "counts": counts,
                "measurements": measurements,
                "camera_frames": self.camera_frames,
                "captured_frames": captured_frames,
                "processed_frames": processed_frames,
                "recorded_frames": recorded_frames,
                "recording_ready": bool(self.recording_path)
                or bool(pending_sample and pending_sample.get("recording_path")),
                "sample_ready": self.sample_ready,
                "inference_busy": self.inference_busy,
                "model_loading": self.model_loading,
                "last_error": self.last_error,
                "completed_samples": len(self.completed_samples),
            }


def get_live_session_state() -> LiveSessionState:
    """Crea un estado independiente para cada sesión del navegador."""
    state = st.session_state.get("live_session_state")
    if isinstance(state, LiveSessionState):
        return state

    # Streamlit puede volver a ejecutar el script y redefinir la clase sin
    # invalidar el objeto que ya vive en session_state. No debemos perder una
    # muestra detenida solo porque su clase proviene de un rerun anterior.
    compatible_state = state is not None and all(
        callable(getattr(state, method_name, None))
        for method_name in (
            "snapshot",
            "set_detection_active",
            "reset_metrics",
            "discard_pending_sample",
            "save_current_sample",
            "completed_samples_snapshot",
        )
    )
    if not compatible_state:
        state = LiveSessionState()
        persisted_samples = st.session_state.get("live_completed_samples")
        if isinstance(persisted_samples, list):
            state.completed_samples = deepcopy(persisted_samples)
        pending_sample = st.session_state.get("live_pending_sample")
        if isinstance(pending_sample, dict) and pending_sample.get("captured_frames", 0) > 0:
            state.pending_sample = deepcopy(pending_sample)
            state.sample_ready = True
            state.captured_frames = int(pending_sample.get("captured_frames", 0))
            state.processed_frames = int(pending_sample.get("processed_frames", 0))
            state.recorded_frames = int(pending_sample.get("recorded_frames", 0))
        st.session_state["live_session_state"] = state
    if not hasattr(state, "measurement_track_ids"):
        # Compatibilidad con una sesión creada antes de incorporar el
        # acumulado de métricas por ID de célula.
        state.measurement_track_ids = set()
    return state


def persist_completed_live_samples(state: LiveSessionState) -> None:
    """Conserva una copia serializable para que un rerun no borre resultados."""
    st.session_state["live_completed_samples"] = state.completed_samples_snapshot()


def persist_pending_live_sample(state: LiveSessionState) -> None:
    """Conserva la muestra detenida aunque Streamlit reconstruya el estado."""
    with state.lock:
        pending_sample = deepcopy(state.pending_sample) if state.sample_ready else None
    if pending_sample is None:
        st.session_state.pop("live_pending_sample", None)
    else:
        st.session_state["live_pending_sample"] = pending_sample


def get_analysis_history() -> list[dict[str, Any]]:
    """Obtiene el historial de la sesión sin guardar videos completos en RAM."""
    history = st.session_state.get(HISTORY_STATE_KEY)
    if not isinstance(history, list):
        history = []
        st.session_state[HISTORY_STATE_KEY] = history
    return history


def _video_path_is_ready(path: str | Path) -> bool:
    """Indica si una ruta apunta a un video que ya terminó de escribirse."""
    try:
        video_path = Path(path)
        return video_path.is_file() and video_path.stat().st_size > 0
    except (OSError, TypeError, ValueError):
        return False


def _transcode_video_for_browser(source_path: Path, destination_path: Path) -> bool:
    """Convierte el MP4 temporal a H.264 para que el reproductor HTML5 lo abra."""
    try:
        import av
    except ImportError:
        return False

    input_container = None
    output_container = None
    success = False
    try:
        input_container = av.open(str(source_path))
        input_stream = next(
            (stream for stream in input_container.streams if stream.type == "video"),
            None,
        )
        if input_stream is None:
            return False

        width = int(input_stream.codec_context.width or 0)
        height = int(input_stream.codec_context.height or 0)
        if width <= 0 or height <= 0:
            return False
        # yuv420p requiere dimensiones pares y el video de la cámara normalmente
        # ya las tiene. El ajuste también cubre cámaras con dimensiones impares.
        width -= width % 2
        height -= height % 2
        if width <= 0 or height <= 0:
            return False

        source_rate = input_stream.average_rate or input_stream.base_rate or 15
        frame_rate = max(1, min(60, round(float(source_rate))))
        output_container = av.open(str(destination_path), mode="w", format="mp4")
        output_stream = None
        for codec_name in ("libx264", "h264", "mpeg4"):
            try:
                output_stream = output_container.add_stream(codec_name, rate=frame_rate)
                break
            except Exception:
                output_stream = None
        if output_stream is None:
            return False

        output_stream.width = width
        output_stream.height = height
        output_stream.pix_fmt = "yuv420p"
        frame_number = 0
        for packet in input_container.demux(input_stream):
            for frame in packet.decode():
                converted = frame.reformat(
                    format="yuv420p",
                    width=width,
                    height=height,
                )
                converted.pts = frame_number
                converted.time_base = Fraction(1, frame_rate)
                for encoded_packet in output_stream.encode(converted):
                    output_container.mux(encoded_packet)
                frame_number += 1

        for encoded_packet in output_stream.encode():
            output_container.mux(encoded_packet)
        success = frame_number > 0
        return success
    except Exception:
        return False
    finally:
        if input_container is not None:
            try:
                input_container.close()
            except Exception:
                pass
        if output_container is not None:
            try:
                output_container.close()
            except Exception:
                pass
        if not success:
            try:
                destination_path.unlink(missing_ok=True)
            except OSError:
                pass


def persist_history_video(source_path: str | Path, history_id: str) -> str:
    """Copia un video guardado a una ruta estable y apta para reproducirse."""
    source = Path(source_path) if str(source_path).strip() else Path()
    if not _video_path_is_ready(source):
        return str(source_path or "")

    try:
        HISTORY_VIDEO_ROOT.mkdir(parents=True, exist_ok=True)
        video_key = hashlib.sha256(history_id.encode("utf-8")).hexdigest()[:24]
        destination = HISTORY_VIDEO_ROOT / f"{video_key}.mp4"
        if _video_path_is_ready(destination):
            return str(destination)

        if not _transcode_video_for_browser(source, destination):
            shutil.copy2(source, destination)
        return str(destination) if _video_path_is_ready(destination) else str(source)
    except OSError:
        return str(source)


def history_video_path(record: dict[str, Any]) -> Path:
    """Obtiene el video del historial y migra registros antiguos si aún existen."""
    recording_value = str(record.get("recording_path") or "").strip()
    source_value = str(record.get("source_recording_path") or "").strip()
    recording_path = Path(recording_value) if recording_value else Path()
    source_path = Path(source_value) if source_value else recording_path
    history_id = str(record.get("history_id") or "history-video")

    if _video_path_is_ready(recording_path):
        # Los registros creados antes de esta mejora apuntan directamente al
        # archivo temporal. Se copian una sola vez a la carpeta del historial.
        if recording_path.parent != HISTORY_VIDEO_ROOT:
            stable_path = persist_history_video(recording_path, history_id)
            if stable_path and stable_path != recording_value:
                record["source_recording_path"] = recording_value
                record["recording_path"] = stable_path
                return Path(stable_path)
        return recording_path

    if _video_path_is_ready(source_path):
        stable_path = persist_history_video(source_path, history_id)
        if stable_path:
            record["source_recording_path"] = str(source_path)
            record["recording_path"] = stable_path
            return Path(stable_path)
    return recording_path


@st.cache_data(show_spinner=False, max_entries=96)
def history_video_thumbnail(video_path: str, modified_ns: int) -> bytes:
    """Extrae una miniatura pequeña del video para la lista del historial."""
    del modified_ns  # Se conserva como parte de la clave de caché para invalidarla al cambiar el archivo.
    capture = cv2.VideoCapture(video_path)
    try:
        if not capture.isOpened():
            return b""
        capture.set(cv2.CAP_PROP_POS_MSEC, 500)
        success, frame = capture.read()
        if not success:
            capture.set(cv2.CAP_PROP_POS_MSEC, 0)
            success, frame = capture.read()
        if not success or frame is None:
            return b""
        frame_height, frame_width = frame.shape[:2]
        if frame_width > 480:
            preview_width = 480
            preview_height = max(1, round(frame_height * preview_width / frame_width))
            frame = cv2.resize(
                frame,
                (preview_width, preview_height),
                interpolation=cv2.INTER_AREA,
            )
        encoded, image_bytes = cv2.imencode(
            ".jpg",
            frame,
            [int(cv2.IMWRITE_JPEG_QUALITY), 78],
        )
        return image_bytes.tobytes() if encoded else b""
    finally:
        capture.release()


def add_live_sample_to_history(
    sample: dict[str, Any],
    low_limit: float = 10.0,
    medium_limit: float = 30.0,
    high_limit: float = 60.0,
) -> None:
    """Registra una muestra guardada para consultarla desde Historial.

    El video se copia a disco, no a la RAM, y queda en una ruta estable durante
    la sesión para que el reproductor pueda abrirlo después de cada rerun.
    """
    history = get_analysis_history()
    source_recording_path = str(sample.get("recording_path") or "").strip()
    sample_number = int(sample.get("sample_number", 0) or 0)
    lot_name = str(sample.get("lot_name") or "Lote sin nombre").strip()
    code = str(sample.get("code") or f"Muestra-{sample_number}").strip()
    existing = next(
        (
            record
            for record in history
            if record.get("analysis_type") == WHITE_SPOT_LABEL
            and (
                record.get("source_recording_path") == source_recording_path
                or record.get("recording_path") == source_recording_path
            )
            and record.get("sample_number") == sample_number
            and record.get("lot_name") == lot_name
        ),
        None,
    )
    affected_percentage, grade, description = sample_result_grade(
        sample.get("counts", {}),
        low_limit,
        medium_limit,
        high_limit,
    )
    if existing is not None:
        if source_recording_path:
            existing["recording_path"] = persist_history_video(
                source_recording_path,
                str(existing.get("history_id") or "history-video"),
            )
            existing["source_recording_path"] = source_recording_path
        existing.update(
            {
                "sample_code": code,
                "grade": grade,
                "description": description,
                "affected_percentage": affected_percentage,
            }
        )
        return

    history_id = f"wssv-{len(history) + 1}-{time.time_ns()}"
    recording_path = persist_history_video(source_recording_path, history_id)
    history.append(
        {
            "history_id": history_id,
            "analysis_type": WHITE_SPOT_LABEL,
            "model_label": WHITE_SPOT_LABEL,
            "sample_number": sample_number,
            "sample_label": str(sample.get("label") or f"Camarón muestra {sample_number}"),
            "sample_code": code,
            "lot_name": lot_name,
            "captured_at": datetime.now().astimezone().strftime("%d/%m/%Y · %H:%M"),
            "grade": grade,
            "description": description,
            "affected_percentage": affected_percentage,
            "counts": dict(sample.get("counts") or {}),
            "measurements": dict(sample.get("measurements") or {})
            if sample.get("measurements")
            else None,
            "recording_path": recording_path,
            "source_recording_path": source_recording_path,
        }
    )


def sync_history_sample_codes(samples: list[dict[str, Any]]) -> None:
    """Refleja en el historial las correcciones de código hechas por el usuario."""
    history = get_analysis_history()
    for sample in samples:
        for record in history:
            if (
                record.get("analysis_type") == WHITE_SPOT_LABEL
                and record.get("sample_number") == sample.get("sample_number")
                and record.get("lot_name") == sample.get("lot_name")
                and (
                    record.get("source_recording_path") == str(sample.get("recording_path") or "")
                    or record.get("recording_path") == str(sample.get("recording_path") or "")
                )
            ):
                record["sample_code"] = str(sample.get("code") or "").strip()


def render_analysis_history() -> None:
    """Renderiza el historial filtrable con reproducción de videos de la sesión."""
    st.subheader("Historial")
    type_column, grade_column = st.columns([1.5, 1.0], gap="medium")
    with type_column:
        selected_type = st.selectbox(
            "Tipo de historial",
            [WHITE_SPOT_LABEL, VACUOLIZATION_LABEL],
            key="history_model_filter",
        )
    records = [
        record
        for record in reversed(get_analysis_history())
        if record.get("model_label") == selected_type
    ]
    grade_options = ["Todos los grados", *(f"Grado {grade}" for grade in range(5)), "Sin datos"]
    with grade_column:
        selected_grade = st.selectbox(
            "Filtrar por grado",
            grade_options,
            key="history_grade_filter",
        )
    if selected_grade != "Todos los grados":
        records = [
            record
            for record in records
            if str(record.get("grade") or "Sin datos").strip() == selected_grade
        ]

    if not records:
        st.markdown(
            f'<div class="history-empty"><strong>{escape(selected_type)}</strong>'
            f"<span>{'No hay muestras de ese grado.' if selected_grade != 'Todos los grados' else 'Aún no hay muestras guardadas en este historial.'}</span></div>",
            unsafe_allow_html=True,
        )
        return

    st.markdown(
        f'<div class="history-count">{len(records)} muestra(s)</div>',
        unsafe_allow_html=True,
    )
    page_size = 12
    page_count = math.ceil(len(records) / page_size)
    if page_count > 1:
        page_state_key = "history_page_" + hashlib.sha256(
            f"{selected_type}|{selected_grade}".encode("utf-8")
        ).hexdigest()[:10]
        page_number = st.selectbox(
            "Página",
            range(1, page_count + 1),
            format_func=lambda page: f"Página {page} de {page_count}",
            key=page_state_key,
        )
    else:
        page_number = 1
    start_index = (page_number - 1) * page_size
    page_records = records[start_index : start_index + page_size]
    for index, record in enumerate(page_records):
        video_path = history_video_path(record)
        video_available = _video_path_is_ready(video_path)
        history_id = str(record.get("history_id") or f"history-{index}")
        record_key = hashlib.sha256(history_id.encode("utf-8")).hexdigest()[:12]
        preview_key = f"history_preview_{record_key}"
        with st.container(key=f"history_row_{record_key}"):
            video_column, metadata_column = st.columns(
                [1.05, 2.8], gap="medium", vertical_alignment="center"
            )
            with video_column:
                with st.container(key=preview_key):
                    if not video_available:
                        st.markdown(
                            '<div class="history-video-missing">Video no disponible</div>',
                            unsafe_allow_html=True,
                        )
                    elif st.session_state.get("history_playing_video") == history_id:
                        st.video(str(video_path), format="video/mp4")
                    else:
                        try:
                            modified_ns = video_path.stat().st_mtime_ns
                        except OSError:
                            modified_ns = 0
                        thumbnail = history_video_thumbnail(str(video_path), modified_ns)
                        if thumbnail:
                            st.image(thumbnail, use_container_width=True)
                            if st.button(
                                "▶",
                                key=f"history_play_{record_key}",
                                help="Reproducir video",
                                type="secondary",
                            ):
                                st.session_state["history_playing_video"] = history_id
                                st.rerun()
                        else:
                            st.markdown(
                                '<div class="history-video-missing">No se pudo cargar el video</div>',
                                unsafe_allow_html=True,
                            )
            with metadata_column:
                sample_code = str(record.get("sample_code") or "—")
                grade = str(record.get("grade") or "Sin datos")
                captured_at = str(record.get("captured_at") or "—")
                lot_name = str(record.get("lot_name") or "—")
                st.markdown(
                    f"""
                    <div class="history-row-meta">
                      <div class="history-row-top">
                        <strong class="history-row-code">{escape(sample_code)}</strong>
                        <span class="history-card-grade">{escape(grade)}</span>
                      </div>
                      <div class="history-row-details">
                        <span><small>Fecha</small><strong>{escape(captured_at)}</strong></span>
                        <span><small>Piscina / lote</small><strong>{escape(lot_name)}</strong></span>
                      </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
        if index < len(page_records) - 1:
            st.markdown('<div class="history-row-divider"></div>', unsafe_allow_html=True)


def normalize_pool_name(value: str) -> str:
    """Normaliza el nombre de una piscina para evitar duplicados accidentales."""
    return " ".join(str(value).strip().split())


def _unique_pool_names(values: list[Any]) -> list[str]:
    """Conserva los nombres de piscina únicos respetando el orden de registro."""
    unique: list[str] = []
    seen: set[str] = set()
    for value in values:
        name = normalize_pool_name(str(value))
        key = name.casefold()
        if name and key not in seen:
            seen.add(key)
            unique.append(name)
    return unique


def get_saved_pools() -> list[str]:
    """Obtiene las piscinas guardadas en la sesión y en la URL del navegador.

    Streamlit Community Cloud no ofrece un disco persistente para datos de usuario.
    La lista se conserva en ``st.session_state`` y también en un parámetro JSON de
    la URL, de modo que vuelve a aparecer al recargar este enlace en el mismo
    navegador. Más adelante puede sustituirse por una base de datos corporativa.
    """
    pools = st.session_state.get("saved_pools")
    if isinstance(pools, list):
        return pools

    raw_value: Any = ""
    try:
        raw_value = st.query_params.get(POOL_QUERY_KEY, "")
    except Exception:
        raw_value = ""
    if isinstance(raw_value, list):
        raw_value = raw_value[-1] if raw_value else ""

    loaded: list[Any] = []
    if raw_value:
        try:
            decoded = json.loads(str(raw_value))
            if isinstance(decoded, list):
                loaded = decoded
        except (TypeError, ValueError, json.JSONDecodeError):
            loaded = []

    pools = _unique_pool_names(loaded)
    st.session_state["saved_pools"] = pools
    return pools


def persist_saved_pools(pools: list[str]) -> None:
    """Guarda la lista de piscinas en la URL cuando el navegador lo permite."""
    try:
        if pools:
            st.query_params[POOL_QUERY_KEY] = json.dumps(pools, ensure_ascii=False)
        elif POOL_QUERY_KEY in st.query_params:
            del st.query_params[POOL_QUERY_KEY]
    except Exception:
        # La lista de la sesión sigue funcionando aunque una versión antigua de
        # Streamlit no exponga query_params.
        pass


def remember_pool(name: str) -> str:
    """Registra una piscina nueva y devuelve el nombre normalizado."""
    normalized = normalize_pool_name(name)
    if not normalized:
        return ""

    pools = get_saved_pools()
    for existing in pools:
        if existing.casefold() == normalized.casefold():
            return existing

    pools.append(normalized)
    st.session_state["saved_pools"] = pools
    persist_saved_pools(pools)
    return normalized


def collect_measurement_samples(
    result: Any,
    model_names: dict[int, str],
    samples: dict[str, dict[str, list[float]]],
    seen_track_ids: set[int] | None = None,
) -> None:
    """Añade métricas sin duplicar una célula rastreada durante la placa."""
    if result.masks is None or result.boxes is None:
        return
    class_ids = result.boxes.cls.int().cpu().tolist()
    track_ids: list[int | None] = [None] * len(class_ids)
    if result.boxes.id is not None:
        tracked_ids = result.boxes.id.int().cpu().tolist()
        track_ids = [int(track_id) for track_id in tracked_ids]

    original_image = getattr(result, "orig_img", None)
    if original_image is not None:
        frame_shape = np.asarray(original_image).shape
    else:
        mask_shape = getattr(result.masks, "orig_shape", None)
        frame_shape = (*mask_shape, 3) if mask_shape else None
    raster_masks = (
        _mask_data_at_frame_size(result, frame_shape)
        if frame_shape is not None
        else None
    )

    if raster_masks is not None and len(raster_masks) == len(class_ids):
        observations = (
            (index, class_id, mask, None)
            for index, (class_id, mask) in enumerate(zip(class_ids, raster_masks))
        )
    else:
        observations = (
            (index, class_id, None, polygon)
            for index, (polygon, class_id) in enumerate(zip(result.masks.xy, class_ids))
        )

    for index, class_id, raster_mask, polygon in observations:
        track_id = track_ids[index] if index < len(track_ids) else None
        if seen_track_ids is not None and track_id is not None and track_id in seen_track_ids:
            continue
        class_name = model_names.get(int(class_id), str(class_id))
        category = normalize_class_name(class_name)
        if category not in {"sana", "enferma"}:
            continue
        if raster_mask is not None:
            binary_mask = np.asarray(raster_mask, dtype=np.uint8)
            area_px2 = float(np.count_nonzero(binary_mask))
            contours, _ = cv2.findContours(
                binary_mask,
                cv2.RETR_EXTERNAL,
                cv2.CHAIN_APPROX_SIMPLE,
            )
            perimeter_px = sum(float(cv2.arcLength(contour, True)) for contour in contours)
        else:
            contour = np.asarray(polygon, dtype=np.float32)
            if contour.shape[0] < 3:
                continue
            area_px2 = float(cv2.contourArea(contour))
            perimeter_px = float(cv2.arcLength(contour, True))
        if area_px2 <= 0 or perimeter_px <= 0:
            continue
        area_values = samples["areas_px2"][category]
        perimeter_values = samples["perimeters_px"][category]
        diameter_values = samples["diameters_px"][category]
        if len(area_values) >= MAX_MEASUREMENT_OBSERVATIONS:
            continue
        area_values.append(area_px2)
        perimeter_values.append(perimeter_px)
        diameter_values.append(math.sqrt(4.0 * area_px2 / math.pi))
        if seen_track_ids is not None and track_id is not None:
            seen_track_ids.add(track_id)


def summarize_measurement_samples(
    samples: dict[str, dict[str, list[float]]], calibration: dict[str, Any]
) -> dict[str, Any] | None:
    """Convierte las métricas en píxeles a promedios en micras."""
    areas_px2 = samples["areas_px2"]
    perimeters_px = samples["perimeters_px"]
    diameters_px = samples["diameters_px"]
    measured_cells = sum(len(areas) for areas in areas_px2.values())
    if not measured_cells:
        return None

    micrometers_per_pixel = float(calibration["micrometers_per_pixel"])
    measurements: dict[str, Any] = {**calibration, "measured_cells": measured_cells}
    for category in ("sana", "enferma"):
        count = len(areas_px2[category])
        measurements[f"{category}_measured_cells"] = count
        measurements[f"{category}_mean_area_um2"] = (
            float(np.mean(areas_px2[category])) * micrometers_per_pixel**2 if count else None
        )
        measurements[f"{category}_mean_perimeter_um"] = (
            float(np.mean(perimeters_px[category])) * micrometers_per_pixel if count else None
        )
        measurements[f"{category}_mean_equivalent_diameter_um"] = (
            float(np.mean(diameters_px[category])) * micrometers_per_pixel if count else None
        )
    return measurements


def measure_cells_from_scale_bar(
    result: Any, model_names: dict[int, str], image: np.ndarray
) -> dict[str, Any] | None:
    """Calcula área y diámetro con barra directa o calibración estimada a 40×."""
    samples = new_measurement_samples()
    collect_measurement_samples(result, model_names, samples)
    return summarize_measurement_samples(samples, calibration_from_image(image))


def grade_from_percentage(
    affected_percentage: float, low_limit: float, medium_limit: float, high_limit: float
) -> tuple[str, str]:
    """Escala provisional configurable hasta contar con la rúbrica validada."""
    if affected_percentage <= 0:
        return "Grado 0", "Sin afectación"
    if affected_percentage <= low_limit:
        return "Grado 1", "Afectación baja"
    if affected_percentage <= medium_limit:
        return "Grado 2", "Afectación moderada"
    if affected_percentage <= high_limit:
        return "Grado 3", "Afectación alta"
    return "Grado 4", "Afectación muy alta"


def format_class_measurement(measurements: dict[str, Any], category: str) -> str:
    """Forma una lectura breve del promedio de una clase de células."""
    count = int(measurements[f"{category}_measured_cells"])
    if count == 0:
        return "No detectadas"
    diameter = float(measurements[f"{category}_mean_equivalent_diameter_um"])
    perimeter = float(measurements[f"{category}_mean_perimeter_um"])
    area = float(measurements[f"{category}_mean_area_um2"])
    return f"Diámetro: {diameter:.2f} µm · Perímetro: {perimeter:.2f} µm · Área: {area:.2f} µm²"


def build_measurement_card(measurements: dict[str, Any], category: str, label: str) -> str:
    """Construye una ficha compacta de perímetro y área para una clase."""
    count = int(measurements[f"{category}_measured_cells"])
    if count == 0:
        details = '<div class="measurement-empty">No detectadas</div>'
    else:
        diameter = float(measurements[f"{category}_mean_equivalent_diameter_um"])
        perimeter = float(measurements[f"{category}_mean_perimeter_um"])
        area = float(measurements[f"{category}_mean_area_um2"])
        details = (
            f'<div class="measurement-row"><span>Diámetro promedio</span><strong>{diameter:.2f} µm</strong></div>'
            f'<div class="measurement-row"><span>Perímetro promedio</span><strong>{perimeter:.2f} µm</strong></div>'
            f'<div class="measurement-row"><span>Área promedio</span><strong>{area:.2f} µm²</strong></div>'
        )
    return (
        f'<article class="result-card measurement-card">'
        f'<span class="measurement-heading">Medidas · {label}</span>'
        f'{details}'
        f'</article>'
    )


def make_report_csv(
    counts: dict[str, int],
    affected_percentage: float,
    grade: str,
    description: str,
    total_label: str,
    measurements: dict[str, Any] | None = None,
) -> bytes:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["Métrica", "Valor"])
    writer.writerow(["Células sanas", counts["sana"]])
    writer.writerow(["Células enfermas", counts["enferma"]])
    writer.writerow(["Células de otras clases", counts["otra"]])
    writer.writerow([total_label, counts["total"]])
    writer.writerow(["Porcentaje de células enfermas", f"{affected_percentage:.2f}%"])
    writer.writerow(["Calificación", grade])
    writer.writerow(["Descripción", description])
    if measurements:
        writer.writerow(["Magnificación de la calibración", "40×"])
        writer.writerow(["Método de calibración", measurements["calibration_method"]])
        writer.writerow(["Referencia de escala", f"{SCALE_BAR_LENGTH_UM:.1f} µm"])
        writer.writerow(["Longitud de referencia", f"{measurements['scale_bar_length_px']:.0f} píxeles"])
        writer.writerow(["Escala", f"{measurements['micrometers_per_pixel']:.6f} µm/píxel"])
        writer.writerow(["Células usadas para medición", measurements["measured_cells"]])
        for category, label in (("sana", "Células sanas"), ("enferma", "Células enfermas")):
            count = int(measurements[f"{category}_measured_cells"])
            writer.writerow([f"{label} usadas para medición", count])
            if count:
                mean_area = float(measurements[f"{category}_mean_area_um2"])
                mean_perimeter = float(measurements[f"{category}_mean_perimeter_um"])
                mean_diameter = float(measurements[f"{category}_mean_equivalent_diameter_um"])
                writer.writerow([f"{label}: área promedio", f"{mean_area:.2f} µm²"])
                writer.writerow(
                    [f"{label}: perímetro promedio", f"{mean_perimeter:.2f} µm"]
                )
                writer.writerow(
                    [f"{label}: diámetro promedio equivalente", f"{mean_diameter:.2f} µm"]
                )
    return output.getvalue().encode("utf-8-sig")


def aggregate_live_sample_metrics(
    samples: list[dict[str, Any]],
) -> tuple[dict[str, int], dict[str, Any] | None]:
    """Suma los conteos y calcula promedios ponderados para todo el lote."""
    counts = {"sana": 0, "enferma": 0, "otra": 0, "total": 0}
    measurement_totals: dict[str, dict[str, float]] = {
        category: {"count": 0.0, "area": 0.0, "perimeter": 0.0, "diameter": 0.0}
        for category in ("sana", "enferma")
    }

    for sample in samples:
        sample_counts = sample["counts"]
        for key in counts:
            counts[key] += int(sample_counts.get(key, 0))

        measurements = sample.get("measurements")
        if not measurements:
            continue
        for category in ("sana", "enferma"):
            measured_count = int(measurements.get(f"{category}_measured_cells", 0) or 0)
            if measured_count <= 0:
                continue
            bucket = measurement_totals[category]
            bucket["count"] += measured_count
            bucket["area"] += float(measurements[f"{category}_mean_area_um2"]) * measured_count
            bucket["perimeter"] += float(measurements[f"{category}_mean_perimeter_um"]) * measured_count
            bucket["diameter"] += float(
                measurements[f"{category}_mean_equivalent_diameter_um"]
            ) * measured_count

    measured_cells = int(sum(bucket["count"] for bucket in measurement_totals.values()))
    if not measured_cells:
        return counts, None

    aggregate_measurements: dict[str, Any] = {
        "calibration_method": "Promedio ponderado de las muestras",
        "measured_cells": measured_cells,
    }
    for category, bucket in measurement_totals.items():
        measured_count = int(bucket["count"])
        aggregate_measurements[f"{category}_measured_cells"] = measured_count
        aggregate_measurements[f"{category}_mean_area_um2"] = (
            bucket["area"] / measured_count if measured_count else None
        )
        aggregate_measurements[f"{category}_mean_perimeter_um"] = (
            bucket["perimeter"] / measured_count if measured_count else None
        )
        aggregate_measurements[f"{category}_mean_equivalent_diameter_um"] = (
            bucket["diameter"] / measured_count if measured_count else None
        )
    return counts, aggregate_measurements


def sample_result_grade(
    counts: dict[str, int], low_limit: float, medium_limit: float, high_limit: float
) -> tuple[float, str, str]:
    """Obtiene afectación y calificación para una muestra o para el lote."""
    denominator = counts["sana"] + counts["enferma"]
    affected_percentage = counts["enferma"] / denominator * 100 if denominator else 0.0
    if denominator:
        grade, description = grade_from_percentage(
            affected_percentage, low_limit, medium_limit, high_limit
        )
    else:
        grade, description = "Sin datos", "No se detectaron células clasificadas"
    return affected_percentage, grade, description


def _append_rows_to_workbook(
    workbook_bytes: bytes | None,
    rows: list[tuple[str, str, str]],
) -> bytes:
    """Añade filas al libro piloto y conserva su tabla de Excel."""
    if not OPENPYXL_AVAILABLE:
        raise OneDriveUploadError(
            "No está disponible openpyxl en el despliegue: "
            f"{OPENPYXL_IMPORT_ERROR or 'dependencia ausente'}."
        )

    if workbook_bytes:
        try:
            workbook = load_workbook(BytesIO(workbook_bytes))
        except Exception as error:
            raise OneDriveUploadError(f"No se pudo abrir el libro de OneDrive: {error}") from error
    else:
        workbook = Workbook()

    worksheet = (
        workbook["Analisis camarones"]
        if "Analisis camarones" in workbook.sheetnames
        else workbook.active
    )
    worksheet.title = "Analisis camarones"

    header_values = [str(cell.value or "").strip() for cell in worksheet[1][: len(EXCEL_HEADERS)]]
    if not any(header_values):
        for column_index, header in enumerate(EXCEL_HEADERS, start=1):
            worksheet.cell(row=1, column=column_index, value=header)
        header_values = list(EXCEL_HEADERS)

    header_columns = {
        value: index + 1
        for index, value in enumerate(header_values)
        if value
    }
    missing_headers = [header for header in EXCEL_HEADERS if header not in header_columns]
    if missing_headers:
        raise OneDriveUploadError(
            "El Excel no tiene las columnas esperadas: " + ", ".join(missing_headers)
        )

    existing_rows = {
        tuple(
            str(worksheet.cell(row=row_number, column=header_columns[header]).value or "").strip()
            for header in EXCEL_HEADERS
        )
        for row_number in range(2, worksheet.max_row + 1)
    }
    next_row = max(worksheet.max_row + 1, 2)
    rows_to_append = [row_values for row_values in rows if row_values not in existing_rows]
    for row_offset, row_values in enumerate(rows_to_append):
        row_number = next_row + row_offset
        worksheet.cell(row=row_number, column=header_columns[EXCEL_HEADERS[0]], value=row_values[0])
        worksheet.cell(row=row_number, column=header_columns[EXCEL_HEADERS[1]], value=row_values[1])
        worksheet.cell(row=row_number, column=header_columns[EXCEL_HEADERS[2]], value=row_values[2])
        existing_rows.add(row_values)

    tables = list(worksheet.tables.values())
    if tables:
        table = tables[0]
        min_column, min_row, max_column, _ = range_boundaries(table.ref)
        table.ref = (
            f"{get_column_letter(min_column)}{min_row}:"
            f"{get_column_letter(max_column)}{worksheet.max_row}"
        )
    else:
        table = Table(
            displayName="AnalisisCamarones",
            ref=f"A1:{get_column_letter(len(EXCEL_HEADERS))}{worksheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    output = BytesIO()
    try:
        workbook.save(output)
    except Exception as error:
        raise OneDriveUploadError(f"No se pudo preparar el Excel: {error}") from error
    finally:
        workbook.close()
    return output.getvalue()


def _live_excel_rows(
    samples: list[dict[str, Any]],
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> list[tuple[str, str, str]]:
    """Convierte las cuatro muestras a las tres columnas del libro corporativo."""
    rows: list[tuple[str, str, str]] = []
    for sample in samples:
        _, white_spot_grade, _ = sample_result_grade(
            sample.get("counts", {}),
            low_limit,
            medium_limit,
            high_limit,
        )
        rows.append(
            (
                str(sample.get("code") or "").strip(),
                "",
                white_spot_grade,
            )
        )
    return rows


def send_live_samples_to_onedrive(
    samples: list[dict[str, Any]],
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> str:
    """Añade las muestras al Excel sincronizado en OneDrive mediante Graph."""
    user = runtime_setting("MS_ONEDRIVE_USER")
    file_path = runtime_setting("MS_ONEDRIVE_FILE_PATH") or "WSSV_Plantilla_Piloto.xlsx"
    if not user:
        raise OneDriveConfigurationError(
            "Configura MS_ONEDRIVE_USER con el correo o ID del usuario propietario del OneDrive."
        )
    access_token = _get_graph_access_token()
    file_url = _onedrive_file_url(user, file_path)
    existing_workbook = _graph_file_content(access_token, file_url)
    updated_workbook = _append_rows_to_workbook(
        existing_workbook,
        _live_excel_rows(samples, low_limit, medium_limit, high_limit),
    )
    _upload_graph_file(access_token, file_url, updated_workbook)
    return file_path


def live_excel_batch_key(
    lot_name: str,
    samples: list[dict[str, Any]],
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> str:
    """Crea una huella para no enviar dos veces el mismo lote por accidente."""
    payload = {
        "lot_name": lot_name.strip(),
        "rows": _live_excel_rows(samples, low_limit, medium_limit, high_limit),
        "recordings": [str(sample.get("recording_path") or "") for sample in samples],
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def make_live_samples_report_csv(
    lot_name: str,
    samples: list[dict[str, Any]],
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> bytes:
    """Crea un reporte CSV con las cuatro muestras y el total del lote."""
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["Nombre del lote / piscina", lot_name])
    writer.writerow(["Muestras analizadas", f"{len(samples)} de {MAX_LIVE_SAMPLES}"])
    writer.writerow([])
    writer.writerow(
        [
            "Muestra",
            "Código",
            "Fotogramas analizados",
            "Células sanas",
            "Células enfermas",
            "Total contabilizado",
            "Afectación",
            "Calificación",
            "Descripción",
        ]
    )

    for sample in samples:
        counts = sample["counts"]
        affected_percentage, grade, description = sample_result_grade(
            counts, low_limit, medium_limit, high_limit
        )
        writer.writerow(
            [
                sample["label"],
                sample["code"],
                sample["processed_frames"],
                counts["sana"],
                counts["enferma"],
                counts["total"],
                f"{affected_percentage:.2f}%",
                grade,
                description,
            ]
        )

    aggregate_counts, aggregate_measurements = aggregate_live_sample_metrics(samples)
    aggregate_percentage, aggregate_grade, aggregate_description = sample_result_grade(
        aggregate_counts, low_limit, medium_limit, high_limit
    )
    writer.writerow([])
    writer.writerow(["MÉTRICA GENERAL DEL LOTE"])
    writer.writerow(["Células sanas", aggregate_counts["sana"]])
    writer.writerow(["Células enfermas", aggregate_counts["enferma"]])
    writer.writerow(["Total contabilizado", aggregate_counts["total"]])
    writer.writerow(["Afectación general", f"{aggregate_percentage:.2f}%"])
    writer.writerow(["Calificación general", aggregate_grade])
    writer.writerow(["Descripción general", aggregate_description])
    if aggregate_measurements:
        writer.writerow(["Células usadas para medición", aggregate_measurements["measured_cells"]])
        for category, label in (("sana", "Sanas"), ("enferma", "Enfermas")):
            count = int(aggregate_measurements[f"{category}_measured_cells"])
            writer.writerow([f"{label} usadas para medición", count])
            if count:
                writer.writerow(
                    [
                        f"{label}: diámetro promedio",
                        f"{aggregate_measurements[f'{category}_mean_equivalent_diameter_um']:.2f} µm",
                    ]
                )
                writer.writerow(
                    [
                        f"{label}: perímetro promedio",
                        f"{aggregate_measurements[f'{category}_mean_perimeter_um']:.2f} µm",
                    ]
                )
                writer.writerow(
                    [
                        f"{label}: área promedio",
                        f"{aggregate_measurements[f'{category}_mean_area_um2']:.2f} µm²",
                    ]
                )
    return output.getvalue().encode("utf-8-sig")


def live_report_filename(lot_name: str) -> str:
    """Genera un nombre de archivo seguro y legible para el reporte del lote."""
    safe_name = "".join(
        character if character.isalnum() or character in {"-", "_"} else "_"
        for character in lot_name.strip()
    ).strip("_")
    return f"reporte_{safe_name or 'lote'}.csv"


def show_summary(
    counts: dict[str, int],
    low_limit: float,
    medium_limit: float,
    high_limit: float,
    total_label: str,
    sidebar: bool = False,
    measurements: dict[str, Any] | None = None,
) -> tuple[float, str, str]:
    """Muestra un panel compacto y legible de resultados para foto y video."""
    target = st.sidebar if sidebar else st
    denominator = counts["sana"] + counts["enferma"]
    affected_percentage = (counts["enferma"] / denominator * 100) if denominator else 0.0
    if denominator:
        grade, description = grade_from_percentage(affected_percentage, low_limit, medium_limit, high_limit)
        result_note = (
            f"<strong>{escape(description)}.</strong> Se detectaron {counts['enferma']} células enfermas "
            f"de {denominator} células clasificadas."
        )
    else:
        grade, description = "Sin datos", "No se detectaron células clasificadas"
        result_note = "<strong>No hay datos para calificar.</strong> Analiza una foto donde se detecten células."
    measurement_cards = ""
    if measurements:
        measurement_cards = (
            build_measurement_card(measurements, "sana", "células sanas")
            + build_measurement_card(measurements, "enferma", "células enfermas")
        )

    target.markdown(
        f"""<section class="results-panel">
<div class="results-heading">Resultados</div>
<div class="results-grid">
<article class="result-card grade-card grade-featured"><span class="result-label">Calificación</span><strong class="grade-value">{escape(grade)}</strong><span class="grade-detail">{affected_percentage:.1f}% enfermas</span></article>
<article class="result-card"><span class="result-label">Células sanas</span><strong class="result-value">{counts['sana']}</strong></article>
<article class="result-card"><span class="result-label">Células enfermas</span><strong class="result-value">{counts['enferma']}</strong></article>
<article class="result-card"><span class="result-label">{escape(total_label)}</span><strong class="result-value">{counts['total']}</strong></article>
{measurement_cards}
</div>
<p class="results-note">{result_note}</p>
</section>""",
        unsafe_allow_html=True,
    )
    if counts["otra"]:
        target.caption(f"Además se detectaron {counts['otra']} células de otras clases.")
    return affected_percentage, grade, description


def process_image(
    model: YOLO, image_bytes: bytes, confidence: float, image_size: int, mask_opacity: float
) -> tuple[np.ndarray, np.ndarray, np.ndarray, dict[str, int], dict[str, Any] | None]:
    """Segmenta una fotografía sin seguimiento: una inferencia y un conteo directo."""
    image_array = np.frombuffer(image_bytes, dtype=np.uint8)
    image = cv2.imdecode(image_array, cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError("No fue posible leer la imagen. Usa PNG, JPG, JPEG o TIFF.")

    result = model.predict(image, conf=confidence, imgsz=image_size, verbose=False)[0]
    model_names = get_model_names(model)
    healthy_masks = mask_frame(result, image.shape, model_names, {"sana"})
    sick_masks = mask_frame(result, image.shape, model_names, {"enferma"})
    measurements = measure_cells_from_scale_bar(result, model_names, image)
    return image, healthy_masks, sick_masks, summarize_instances(result, model_names), measurements


def process_video(
    model: YOLO,
    uploaded_video: Any,
    confidence: float,
    image_size: int,
    mask_opacity: float,
    frame_stride: int,
    show_masks: bool,
    progress_bar: Any,
    status_text: Any,
    preview_placeholder: Any | None = None,
) -> tuple[bytes, dict[str, int], int, dict[str, Any] | None]:
    """Procesa el video, actualiza una previsualización y cuenta células únicas."""
    suffix = Path(uploaded_video.name).suffix or ".mp4"
    source_file = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    output_file = tempfile.NamedTemporaryFile(suffix=".mp4", delete=False)
    source_path = Path(source_file.name)
    output_path = Path(output_file.name)
    source_file.close()
    output_file.close()

    capture: cv2.VideoCapture | None = None
    writer: cv2.VideoWriter | None = None
    results: Any | None = None
    track_votes: dict[int, Counter] | None = None
    video_measurement_samples: dict[str, dict[str, list[float]]] | None = None
    try:
        uploaded_video.seek(0)
        with source_path.open("wb") as target:
            shutil.copyfileobj(uploaded_video, target)

        capture = cv2.VideoCapture(str(source_path))
        if not capture.isOpened():
            raise ValueError("No fue posible abrir el video. Usa un archivo MP4, AVI o MOV válido.")

        fps = capture.get(cv2.CAP_PROP_FPS)
        fps = fps if fps and fps > 0 else 25.0
        width = int(capture.get(cv2.CAP_PROP_FRAME_WIDTH))
        height = int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT))
        frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
        if width <= 0 or height <= 0:
            raise ValueError("El video no tiene dimensiones válidas.")

        # Solo se usa OpenCV para conocer las propiedades del video; YOLO recibe
        # el archivo completo en una única llamada, no un fotograma por llamada.
        capture.release()
        capture = None

        # Se guardan los fotogramas analizados a una tasa proporcional para que
        # el video final conserve aproximadamente la misma duración.
        output_fps = max(fps / frame_stride, 1.0)
        output_width = min(width, VIDEO_OUTPUT_MAX_WIDTH)
        output_height = max(1, round(height * output_width / width))
        writer = cv2.VideoWriter(
            str(output_path),
            cv2.VideoWriter_fourcc(*"mp4v"),
            output_fps,
            (output_width, output_height),
        )
        if not writer.isOpened():
            raise RuntimeError("No se pudo crear el video de salida en formato MP4.")

        reset_trackers(model)
        model_names = get_model_names(model)
        track_votes = defaultdict(Counter)
        tracked_detections = 0
        processed_frames = 0
        last_preview_update = 0.0
        video_measurement_samples = new_measurement_samples()
        video_measurement_track_ids: set[int] = set()
        video_calibration: dict[str, Any] | None = None

        results = model.track(
            source=str(source_path),
            stream=True,
            persist=True,
            tracker="bytetrack.yaml",
            conf=confidence,
            imgsz=image_size,
            retina_masks=True,
            vid_stride=frame_stride,
            verbose=False,
        )
        for result in results:
            if video_calibration is None:
                video_calibration = calibration_from_image(result.orig_img)
            collect_measurement_samples(
                result,
                model_names,
                video_measurement_samples,
                video_measurement_track_ids,
            )
            masks_frame = mask_frame(result, result.orig_img.shape, model_names)
            output_frame = (
                overlay_masks(result.orig_img, masks_frame, mask_opacity)
                if show_masks
                else result.orig_img
            )
            if output_frame.shape[1] != output_width or output_frame.shape[0] != output_height:
                output_frame = cv2.resize(
                    output_frame,
                    (output_width, output_height),
                    interpolation=cv2.INTER_AREA,
                )
            writer.write(output_frame)
            tracked_detections += update_track_votes(result, track_votes, model_names)
            processed_frames += 1
            source_frame = (
                min(processed_frames * frame_stride, frame_count)
                if frame_count > 0
                else processed_frames * frame_stride
            )

            if frame_count > 0:
                progress_bar.progress(min(source_frame / frame_count, 1.0))
            status_text.caption(f"Procesando fotograma {source_frame:,} de {frame_count:,}…")

            # Se muestra el video durante el análisis, pero se limita a cuatro
            # actualizaciones por segundo para no ralentizar YOLO.
            now = time.monotonic()
            if preview_placeholder is not None and (
                processed_frames == 1 or now - last_preview_update >= 0.25
            ):
                preview_frame = output_frame
                preview_width = preview_frame.shape[1]
                if preview_width > 960:
                    scale = 960 / preview_width
                    preview_frame = cv2.resize(
                        preview_frame,
                        (960, max(1, round(preview_frame.shape[0] * scale))),
                        interpolation=cv2.INTER_AREA,
                    )
                preview_placeholder.image(
                    preview_frame,
                    channels="BGR",
                    caption=(
                        "Video original con máscaras · detección en curso"
                        if show_masks
                        else "Video original · detección en curso"
                    ),
                    use_container_width=True,
                )
                last_preview_update = now

        writer.release()
        writer = None
        progress_bar.progress(1.0)

        if processed_frames == 0:
            raise ValueError("No se encontraron fotogramas para procesar.")

        video_measurements = (
            summarize_measurement_samples(video_measurement_samples, video_calibration)
            if video_calibration is not None
            else None
        )
        return output_path.read_bytes(), summarize_tracks(track_votes), tracked_detections, video_measurements
    finally:
        if capture is not None:
            capture.release()
        if writer is not None:
            writer.release()
        for temporary_path in (source_path, output_path):
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
        results = None
        if track_votes is not None:
            track_votes.clear()
        if video_measurement_samples is not None:
            for values_by_category in video_measurement_samples.values():
                for values in values_by_category.values():
                    values.clear()
        gc.collect()


def render_photo_mode(
    model: YOLO | None,
    confidence: float,
    image_size: int,
    mask_opacity: float,
    low_limit: float,
    medium_limit: float,
    high_limit: float,
    model_loader: Callable[[], YOLO] | None = None,
) -> None:
    """Renderiza el flujo rápido de una sola fotografía."""
    st.markdown(
        """
        <div class="upload-heading">
          <span class="upload-step">1</span>
          <div><strong>Carga una fotografía</strong><small>PNG, JPG o TIFF</small></div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    active_photo = st.session_state.get("active_photo")
    if active_photo is None:
        upload_revision = st.session_state.get("photo_upload_revision", 0)
        uploaded_image = st.file_uploader(
            "Seleccionar fotografía",
            type=["png", "jpg", "jpeg", "tif", "tiff"],
            label_visibility="collapsed",
            key=f"photo_upload_{upload_revision}",
        )
        if uploaded_image is not None:
            st.session_state["active_photo"] = {
                "name": uploaded_image.name,
                "bytes": uploaded_image.getvalue(),
            }
            st.rerun()

        st.button("Analizar foto", type="primary", use_container_width=True, disabled=True)
        st.info("Selecciona una foto para habilitar el botón de análisis.")
        return

    file_name_column, change_photo_column = st.columns([5, 1])
    with file_name_column:
        st.markdown(
            f'<div class="photo-file-row"><span>Foto cargada</span><strong>{escape(active_photo["name"])}</strong></div>',
            unsafe_allow_html=True,
        )
    with change_photo_column:
        change_photo = st.button("Cambiar", key="change_photo", use_container_width=True)
    if change_photo:
        st.session_state.pop("active_photo", None)
        st.session_state.pop("photo_analysis", None)
        st.session_state["photo_upload_revision"] = st.session_state.get("photo_upload_revision", 0) + 1
        st.rerun()

    analyze_photo = st.button("Analizar foto", type="primary", use_container_width=True)
    image_bytes = active_photo["bytes"]
    image_id = hashlib.sha256(image_bytes).hexdigest()
    saved_analysis = st.session_state.get("photo_analysis")
    if saved_analysis and (
        saved_analysis["image_id"] != image_id or saved_analysis.get("render_version") != PHOTO_RENDER_VERSION
    ):
        st.session_state.pop("photo_analysis", None)
        saved_analysis = None

    if not analyze_photo:
        if saved_analysis is None:
            return
    else:
        with st.spinner("Analizando la foto…"):
            try:
                analysis_model = model
                if analysis_model is None:
                    if model_loader is None:
                        raise RuntimeError("El modelo de inferencia aún no está disponible.")
                    analysis_model = model_loader()
                original_image, healthy_masks, sick_masks, counts, measurements = process_image(
                    analysis_model, image_bytes, confidence, image_size, mask_opacity
                )
            except Exception as error:
                st.exception(error)
                return
        st.session_state["photo_analysis"] = {
            "image_id": image_id,
            "render_version": PHOTO_RENDER_VERSION,
            "original_image": original_image,
            "healthy_masks": healthy_masks,
            "sick_masks": sick_masks,
            "counts": counts,
            "measurements": measurements,
        }
        saved_analysis = st.session_state["photo_analysis"]

    affected_percentage, grade, description = show_summary(
        saved_analysis["counts"],
        low_limit,
        medium_limit,
        high_limit,
        "Total de células detectadas",
        sidebar=True,
        measurements=saved_analysis["measurements"],
    )
    visible_photo_masks = st.pills(
        "Controles de máscaras",
        ["Sanas", "Enfermas"],
        selection_mode="multi",
        default=["Sanas", "Enfermas"],
        key=f"photo_visible_masks_{image_id}",
        label_visibility="collapsed",
    )
    show_healthy_masks = "Sanas" in visible_photo_masks
    show_sick_masks = "Enfermas" in visible_photo_masks

    selected_masks = np.zeros_like(saved_analysis["original_image"])
    if show_healthy_masks:
        selected_masks = cv2.bitwise_or(selected_masks, saved_analysis["healthy_masks"])
    if show_sick_masks:
        selected_masks = cv2.bitwise_or(selected_masks, saved_analysis["sick_masks"])
    has_selected_masks = show_healthy_masks or show_sick_masks
    image_to_display = (
        overlay_masks(saved_analysis["original_image"], selected_masks, mask_opacity)
        if has_selected_masks
        else saved_analysis["original_image"]
    )
    legend_items: list[str] = []
    if show_healthy_masks:
        legend_items.append('<span><i class="legend-dot legend-healthy"></i>Célula sana</span>')
    if show_sick_masks:
        legend_items.append('<span><i class="legend-dot legend-sick"></i>Célula enferma</span>')
    if legend_items:
        st.markdown(
            '<div class="mask-legend">' + "".join(legend_items) + "</div>",
            unsafe_allow_html=True,
        )
    st.image(image_to_display, channels="BGR")
    if saved_analysis["measurements"]:
        measurements = saved_analysis["measurements"]
        scale = measurements["micrometers_per_pixel"]
        bar_pixels = measurements["scale_bar_length_px"]
        healthy_measurement = format_class_measurement(measurements, "sana")
        sick_measurement = format_class_measurement(measurements, "enferma")
        st.markdown(
            f"""
            <section class="measurement-inline">
              <span class="measurement-inline-title">Estimación promedio por clase</span>
              <strong>Sanas: {healthy_measurement}</strong>
              <strong>Enfermas: {sick_measurement}</strong>
            </section>
            """,
            unsafe_allow_html=True,
        )
        st.caption(
            f"Calibración {measurements['calibration_method']}: referencia de 50 µm · "
            f"{bar_pixels:.0f} px · escala {scale:.4f} µm/píxel."
        )
    else:
        st.info("No se segmentaron células; por eso no es posible calcular un promedio de tamaño.")

    encoded, image_png = cv2.imencode(".png", image_to_display)
    if encoded:
        download_image, download_report = st.columns(2)
        with download_image:
            st.download_button(
                "Descargar imagen PNG",
                data=image_png.tobytes(),
                file_name="celulas_segmentadas.png" if has_selected_masks else "celulas_original.png",
                mime="image/png",
                use_container_width=True,
            )
        with download_report:
            report = make_report_csv(
                saved_analysis["counts"],
                affected_percentage,
                grade,
                description,
                "Total de células detectadas",
                measurements=saved_analysis["measurements"],
            )
            st.download_button(
                "Descargar reporte CSV",
                data=report,
                file_name="reporte_celular.csv",
                mime="text/csv",
                use_container_width=True,
            )

    st.caption(
        "Cada célula segmentada en la fotografía se cuenta una sola vez. "
        "La calificación debe validarse con tu criterio científico o clínico."
    )


def render_live_metrics_panel(
    state: LiveSessionState,
    lot_name: str,
    current_sample_code: str,
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> None:
    """Muestra las cuatro fichas de camarones y el acumulado del lote."""

    def metric_card(
        label: str,
        code: str,
        counts: dict[str, int],
        measurements: dict[str, Any] | None,
        low_limit: float,
        medium_limit: float,
        high_limit: float,
        card_class: str = "",
    ) -> str:
        affected_percentage, grade, _ = sample_result_grade(
            counts, low_limit, medium_limit, high_limit
        )
        measurement_text = ""
        if measurements:
            measurement_text = (
                f'<div class="live-card-measurements">'
                f"{escape(format_class_measurement(measurements, 'sana'))}<br>"
                f"{escape(format_class_measurement(measurements, 'enferma'))}"
                f"</div>"
            )
        return f"""
        <article class="live-sample-card {card_class}">
          <div class="live-sample-card-heading">
            <strong>{escape(label)}</strong>
            <span>Código: {escape(code)}</span>
          </div>
          <div class="live-card-grid">
            <div><span>Sanas</span><strong>{counts['sana']}</strong></div>
            <div><span>Enfermas</span><strong>{counts['enferma']}</strong></div>
            <div><span>Total</span><strong>{counts['total']}</strong></div>
            <div><span>Afectación</span><strong>{affected_percentage:.1f}%</strong></div>
          </div>
          <div class="live-card-grade"><span>Calificación</span><strong>{escape(grade)}</strong></div>
          {measurement_text}
        </article>
        """

    def render_metrics() -> None:
        snapshot = state.snapshot()
        completed_samples = state.completed_samples_snapshot()
        detection_requested = bool(st.session_state.get("live_detection_requested", False))
        counts = snapshot["counts"]
        sample_number = min(len(completed_samples) + 1, MAX_LIVE_SAMPLES)
        current_card = ""
        if (
            detection_requested
            or snapshot["detection_active"]
            or snapshot["sample_ready"]
            or snapshot["captured_frames"] > 0
        ) and len(completed_samples) < MAX_LIVE_SAMPLES:
            current_card = metric_card(
                f"Camarón muestra {sample_number}",
                current_sample_code or "Código pendiente",
                counts,
                snapshot["measurements"],
                low_limit,
                medium_limit,
                high_limit,
                "live-current-card",
            )

        sample_cards = "".join(
            metric_card(
                sample["label"],
                sample["code"],
                sample["counts"],
                sample.get("measurements"),
                low_limit,
                medium_limit,
                high_limit,
            )
            for sample in completed_samples
        )
        aggregate_counts, aggregate_measurements = aggregate_live_sample_metrics(completed_samples)
        aggregate_card = ""
        if completed_samples:
            aggregate_card = metric_card(
                f"Métrica general del lote ({len(completed_samples)}/{MAX_LIVE_SAMPLES} muestras)",
                "Acumulado",
                aggregate_counts,
                aggregate_measurements,
                low_limit,
                medium_limit,
                high_limit,
                "live-aggregate-card",
            )

        cards = current_card + sample_cards + aggregate_card
        if not cards:
            st.caption("Inicia la detección para comenzar a capturar la muestra.")
        else:
            st.markdown(
                f"""
                <div class="live-sample-cards">{cards}</div>
                """,
                unsafe_allow_html=True,
            )

        if (
            snapshot["captured_frames"] > 0
            or snapshot["sample_ready"]
            or snapshot["detection_active"]
            or detection_requested
        ):
            status = (
                "Grabando muestra"
                if snapshot["detection_active"] or detection_requested
                else "Muestra detenida"
            )
            st.caption(
                f"{status} · {counts['total']:,} células únicas confirmadas · "
                f"{snapshot['captured_frames']:,} fotogramas capturados · "
                f"{snapshot['processed_frames']:,} inferidos"
            )

        if snapshot["model_loading"]:
            st.caption("Preparando el modelo… el video continúa grabándose.")

        if snapshot["last_error"]:
            st.warning(f"La cámara continúa, pero una inferencia dio error: {snapshot['last_error']}")

    # El fragmento refresca solamente el panel, sin reiniciar la cámara completa.
    if hasattr(st, "fragment"):
        @st.fragment(run_every="1s")
        def live_metrics_fragment() -> None:
            render_metrics()

        live_metrics_fragment()
    else:
        render_metrics()


def render_live_sample_results(
    state: LiveSessionState,
    lot_name: str,
    camera_is_playing: bool,
    low_limit: float,
    medium_limit: float,
    high_limit: float,
) -> None:
    """Muestra la tabla final y habilita el reporte al completar cuatro muestras."""
    samples = state.completed_samples_snapshot()
    if not samples:
        return

    rows: list[dict[str, Any]] = []
    for sample in samples:
        affected_percentage, grade, _ = sample_result_grade(
            sample["counts"], low_limit, medium_limit, high_limit
        )
        rows.append(
            {
                "Muestra": sample["label"],
                "Código": sample["code"],
                "Sanas": sample["counts"]["sana"],
                "Enfermas": sample["counts"]["enferma"],
                "Total": sample["counts"]["total"],
                "Afectación": f"{affected_percentage:.1f}%",
                "Calificación": grade,
            }
        )

    st.subheader(f"Muestras de {lot_name or 'lote sin nombre'}")
    st.dataframe(rows, hide_index=True, use_container_width=True)

    aggregate_counts, aggregate_measurements = aggregate_live_sample_metrics(samples)
    aggregate_percentage, aggregate_grade, aggregate_description = sample_result_grade(
        aggregate_counts, low_limit, medium_limit, high_limit
    )
    st.markdown(
        f"""
        <section class="live-overall-summary">
          <strong>Métrica general del lote</strong>
          <span>{aggregate_counts['sana']} sanas · {aggregate_counts['enferma']} enfermas · {aggregate_counts['total']} células en total</span>
          <span>{aggregate_percentage:.1f}% de afectación · {escape(aggregate_grade)} · {escape(aggregate_description)}</span>
        </section>
        """,
        unsafe_allow_html=True,
    )

    with st.expander("Editar códigos guardados", expanded=False):
        edited_codes: dict[int, str] = {}
        with st.form("live_edit_sample_codes_form", clear_on_submit=False):
            edit_columns = st.columns(2, gap="small")
            for index, sample in enumerate(samples):
                sample_number = int(sample["sample_number"])
                field_key = f"live_edit_sample_code_{sample_number}"
                if field_key not in st.session_state:
                    st.session_state[field_key] = sample["code"]
                with edit_columns[index % 2]:
                    edited_codes[sample_number] = st.text_input(
                        sample["label"],
                        key=field_key,
                    )
            save_code_edits = st.form_submit_button(
                "Guardar cambios",
                use_container_width=True,
            )
        if save_code_edits:
            for sample_number, code in edited_codes.items():
                state.update_completed_sample_code(sample_number, code)
            persist_completed_live_samples(state)
            sync_history_sample_codes(state.completed_samples_snapshot())
            st.rerun()

    if len(samples) == MAX_LIVE_SAMPLES:
        report = make_live_samples_report_csv(
            lot_name or "Lote sin nombre",
            samples,
            low_limit,
            medium_limit,
            high_limit,
        )
        st.download_button(
            "Descargar reporte del lote (4 muestras)",
            data=report,
            file_name=live_report_filename(lot_name),
            mime="text/csv",
            use_container_width=True,
            key="live_samples_report",
        )

        excel_preview = [
            {
                "Código": row[0],
                "Vacuolización": row[1],
                "White Spot (WSSV)": row[2],
            }
            for row in _live_excel_rows(samples, low_limit, medium_limit, high_limit)
        ]
        st.dataframe(excel_preview, hide_index=True, use_container_width=True)
        batch_key = live_excel_batch_key(
            lot_name or "Lote sin nombre",
            samples,
            low_limit,
            medium_limit,
            high_limit,
        )
        sent_batches = st.session_state.get(ONEDRIVE_SENT_BATCHES_KEY, [])
        if not isinstance(sent_batches, list):
            sent_batches = []
        if batch_key in sent_batches:
            st.success("Este lote ya fue enviado al Excel de OneDrive.")
        elif st.button(
            "Enviar al Excel",
            type="primary",
            use_container_width=True,
            key="live_send_to_onedrive",
        ):
            with st.spinner("Enviando las cuatro muestras a OneDrive…"):
                try:
                    sent_file_path = send_live_samples_to_onedrive(
                        samples,
                        low_limit,
                        medium_limit,
                        high_limit,
                    )
                except OneDriveConfigurationError as error:
                    st.warning(str(error))
                except OneDriveUploadError as error:
                    st.error(str(error))
                else:
                    sent_batches.append(batch_key)
                    st.session_state[ONEDRIVE_SENT_BATCHES_KEY] = sent_batches
                    st.success(f"Lote enviado a OneDrive: {sent_file_path}")
    if st.button(
        "Nuevo lote / análisis",
        key="live_new_analysis",
        use_container_width=True,
        disabled=camera_is_playing or state.snapshot()["detection_active"],
    ):
        state.set_detection_active(False)
        st.session_state["live_detection_requested"] = False
        state.reset_metrics()
        state.clear_completed_samples()
        st.session_state.pop("live_completed_samples", None)
        st.session_state.pop("live_pending_sample", None)
        st.session_state.pop("live_lot_name", None)
        for key in list(st.session_state):
            if key.startswith(("live_sample_code_", "live_edit_sample_code_")):
                st.session_state.pop(key, None)
        st.session_state["live_camera_playing"] = False
        st.rerun()


def render_live_camera(
    model: YOLO | None,
    confidence: float,
    mask_opacity: float,
    model_loader: Callable[[], YOLO] | None = None,
) -> None:
    """Mantiene la cámara activa y ejecuta el tracker solo cuando el usuario lo inicia."""
    if not WEBRTC_AVAILABLE:
        st.error("No se pudo cargar el componente de cámara en vivo.")
        st.caption(
            "Detalle de la dependencia: "
            f"{WEBRTC_IMPORT_ERROR or 'error de importación desconocido'}"
        )
        return

    state = get_live_session_state()
    save_feedback = st.session_state.pop("live_save_feedback", "")
    if save_feedback:
        st.success(save_feedback)
    model_names = get_model_names(model) if model is not None else None
    with state.lock:
        # El procesador WebRTC puede sobrevivir a un rerun de Streamlit; guardar
        # el modelo en el estado permite que detecte aun si el callback anterior
        # se creó antes de que el usuario pulsara "Iniciar detección".
        if model is not None:
            state.inference_model = model
            state.inference_model_names = model_names or DEFAULT_MODEL_NAMES.copy()
    if "live_camera_playing" not in st.session_state:
        st.session_state["live_camera_playing"] = False
    if "live_camera_requested" not in st.session_state:
        st.session_state["live_camera_requested"] = False
    if "live_detection_requested" not in st.session_state:
        st.session_state["live_detection_requested"] = state.snapshot()["detection_active"]
    if st.session_state["live_detection_requested"] and not state.snapshot()["detection_active"]:
        # El componente WebRTC puede provocar un rerun mientras el callback
        # del botón todavía termina. Recuperar aquí la bandera evita que el
        # primer fotograma vea la detección como detenida.
        state.set_detection_active(True)

    camera_is_playing = bool(st.session_state.get("live_camera_playing", False))
    camera_requested = bool(st.session_state.get("live_camera_requested", False))
    completed_count = len(state.completed_samples_snapshot())
    saved_pools = get_saved_pools()
    if "live_pool_selector_pending" in st.session_state:
        st.session_state["live_pool_selector"] = st.session_state.pop("live_pool_selector_pending")
    if "live_pool_selector" not in st.session_state:
        st.session_state["live_pool_selector"] = (
            saved_pools[0] if saved_pools else NEW_POOL_OPTION
        )
    pool_options = [*saved_pools, NEW_POOL_OPTION]

    pool_column, sample_column = st.columns([1.0, 1.0], gap="large")
    with pool_column:
        selected_pool = st.selectbox(
            "Piscina registrada",
            pool_options,
            key="live_pool_selector",
            disabled=completed_count > 0,
        )
        if selected_pool == NEW_POOL_OPTION:
            new_pool_name = st.text_input(
                "Nombre o código de la nueva piscina",
                key="live_new_pool_name",
                placeholder="Ej.: Piscina Norte 01",
                disabled=completed_count > 0,
            )
            if st.button(
                "Guardar piscina en la lista",
                key="live_save_pool",
                use_container_width=True,
                disabled=completed_count > 0,
            ):
                registered_pool = remember_pool(new_pool_name)
                if registered_pool:
                    st.session_state["live_pool_selector_pending"] = registered_pool
                    st.rerun()
                st.warning("Escribe un nombre o código para guardar la piscina.")
            lot_name = normalize_pool_name(new_pool_name)
        else:
            lot_name = normalize_pool_name(selected_pool)
        st.session_state["live_lot_name"] = lot_name

    current_sample_number = min(completed_count + 1, MAX_LIVE_SAMPLES)
    sample_code_key = f"live_sample_code_{current_sample_number}"
    if sample_code_key not in st.session_state:
        st.session_state[sample_code_key] = f"CAM-{current_sample_number:02d}"
    with sample_column:
        sample_code = st.text_input(
            f"Código del camarón · muestra {current_sample_number} de {MAX_LIVE_SAMPLES}",
            key=sample_code_key,
        ).strip()
    resolution_label = st.selectbox(
        "Resolución de captura de video",
        list(LIVE_RESOLUTIONS),
        index=0,
        key="live_resolution_label",
        disabled=camera_is_playing or camera_requested,
    )
    resolution = LIVE_RESOLUTIONS[resolution_label]

    # Los tres controles viven fuera del video. Se usan las columnas directamente
    # en lugar de `st.empty()`: los placeholders podían conservar un contenedor
    # vacío después de un rerun de WebRTC y aparentar un segundo control blanco.
    camera_action, detection_action, save_action = st.columns(
        [1.05, 1.25, 1.15], gap="small"
    )

    def toggle_live_camera() -> None:
        """Solicita iniciar o detener la cámara desde el botón exterior."""
        requested = bool(st.session_state.get("live_camera_requested", False))
        if requested:
            # Si se detiene la cámara durante una muestra, la muestra queda
            # cerrada pero sus métricas siguen disponibles para guardarlas.
            if state.snapshot()["detection_active"]:
                state.set_detection_active(False)
                persist_pending_live_sample(state)
                st.session_state["live_detection_requested"] = False
            st.session_state["live_camera_requested"] = False
            st.session_state["live_camera_playing"] = False
            return
        st.session_state["live_camera_requested"] = True
        st.session_state.pop("live_camera_start_required", None)

    def toggle_live_detection() -> None:
        """Cambia la detección sin desmontar el componente WebRTC."""
        snapshot = state.snapshot()
        if snapshot["detection_active"] or st.session_state.get(
            "live_detection_requested", False
        ):
            state.set_detection_active(False)
            persist_pending_live_sample(state)
            st.session_state["live_detection_requested"] = False
            return

        # Una muestra detenida debe guardarse antes de comenzar la siguiente.
        # Así las cuatro placas permanecen independientes y nunca se descarta
        # silenciosamente el resultado anterior al iniciar la nueva.
        if snapshot["sample_ready"]:
            return

        state.reset_metrics()
        state.measurement_track_ids = set()
        st.session_state.pop("live_pending_sample", None)
        with state.lock:
            active_model = state.inference_model
        if active_model is not None:
            with LIVE_INFERENCE_LOCK:
                reset_trackers(active_model)
        if lot_name:
            remember_pool(lot_name)
        if not bool(st.session_state.get("live_camera_requested", False)):
            st.session_state["live_detection_requested"] = False
            st.session_state["live_camera_start_required"] = True
            return
        st.session_state["live_detection_requested"] = True
        state.set_detection_active(True)

    def ensure_live_model() -> None:
        """Prepara YOLO en segundo plano sin detener el flujo de la cámara."""
        if model_loader is None:
            return
        with state.lock:
            if state.inference_model is not None or state.model_loading:
                return
            state.model_loading = True

        def load_in_background() -> None:
            try:
                loaded_model = model_loader()
                loaded_names = get_model_names(loaded_model)
                with state.lock:
                    state.inference_model = loaded_model
                    state.inference_model_names = loaded_names
                    state.model_loading = False
                    state.inference_retry_at = 0.0
                    state.last_error = ""
            except Exception as error:
                with state.lock:
                    state.model_loading = False
                    state.last_error = f"{type(error).__name__}: {error}"
                    state.inference_retry_at = time.monotonic() + LIVE_INFERENCE_RETRY_SECONDS

        Thread(target=load_in_background, daemon=True).start()

    def run_live_inference(image: np.ndarray, generation: int) -> None:
        """Ejecuta YOLO fuera del callback para no congelar el video al comenzar."""
        try:
            with state.lock:
                active_model = state.inference_model
                model_names = dict(state.inference_model_names)
            if active_model is None:
                raise RuntimeError("El modelo aún no está listo; vuelve a iniciar la detección.")
            with LIVE_INFERENCE_LOCK:
                results = active_model.track(
                    source=image,
                    conf=confidence,
                    imgsz=resolution["inference_size"],
                    retina_masks=True,
                    persist=True,
                    tracker=str(LIVE_TRACKER_CONFIG),
                    verbose=False,
                )
            if not results:
                empty_mask = np.zeros_like(image)
                with state.lock:
                    if state.inference_generation == generation and state.detection_active:
                        state.cached_shape = image.shape
                        state.cached_healthy_masks = empty_mask
                        state.cached_sick_masks = empty_mask.copy()
                        if state.calibration is None:
                            state.calibration = calibration_from_image(image)
                        state.processed_frames += 1
                        state.last_error = ""
                        state.inference_retry_at = 0.0
                return

            result = results[0]
            healthy_masks = mask_frame(result, image.shape, model_names, {"sana"})
            sick_masks = mask_frame(result, image.shape, model_names, {"enferma"})
            with state.lock:
                # Si el usuario detuvo o reinició la muestra mientras YOLO
                # trabajaba, el resultado viejo no debe contaminar la siguiente.
                if state.inference_generation != generation or not state.detection_active:
                    return
                state.cached_shape = image.shape
                state.cached_healthy_masks = healthy_masks
                state.cached_sick_masks = sick_masks
                if state.calibration is None:
                    state.calibration = calibration_from_image(image)
                update_track_votes(result, state.track_votes, model_names)
                collect_measurement_samples(
                    result,
                    model_names,
                    state.measurement_samples,
                    state.measurement_track_ids,
                )
                state.processed_frames += 1
                state.last_error = ""
                state.inference_retry_at = 0.0
        except Exception as error:
            with state.lock:
                if state.inference_generation == generation:
                    state.last_error = f"{type(error).__name__}: {error}"
                    # Evita crear un hilo nuevo por cada fotograma si el backend
                    # devuelve un error rápido o tarda en preparar el predictor.
                    state.inference_retry_at = time.monotonic() + LIVE_INFERENCE_RETRY_SECONDS
        finally:
            with state.lock:
                if state.inference_generation == generation:
                    state.inference_busy = False

    def process_live_frame(frame: Any) -> Any:
        image: np.ndarray | None = None

        def make_output_frame(output: np.ndarray) -> av.VideoFrame:
            """Devuelve un cuadro reproducible y conserva su reloj WebRTC."""
            output_frame = av.VideoFrame.from_ndarray(
                np.ascontiguousarray(output),
                format="bgr24",
            )
            # aiortc necesita conservar PTS/time_base para que Chrome pueda
            # reproducir el track de salida de forma continua.
            output_frame.pts = getattr(frame, "pts", None)
            output_frame.time_base = getattr(frame, "time_base", None)
            return output_frame

        try:
            image = frame.to_ndarray(format="bgr24")
            with state.lock:
                state.frame_number += 1
                state.camera_frames += 1
                frame_number = state.frame_number
                detection_active = state.detection_active
                cached_shape = state.cached_shape
                cached_healthy_masks = state.cached_healthy_masks
                cached_sick_masks = state.cached_sick_masks
                inference_busy = state.inference_busy
                active_model_available = state.inference_model is not None
                inference_generation = state.inference_generation
                inference_retry_at = state.inference_retry_at

            if not detection_active:
                return make_output_frame(image)

            # La captura se contabiliza antes de esperar a OpenVINO. Así,
            # detener la sesión siempre deja una muestra guardable, incluso
            # si una inferencia puntual tarda o devuelve cero detecciones.
            with state.lock:
                state.captured_frames += 1

            needs_inference = (
                cached_healthy_masks is None
                or cached_sick_masks is None
                or cached_shape != image.shape
                or frame_number % resolution["inference_every"] == 1
            )
            if (
                needs_inference
                and active_model_available
                and not inference_busy
                and time.monotonic() >= inference_retry_at
            ):
                # Solo se agenda un fotograma a la vez. El callback sigue
                # entregando video mientras OpenVINO termina el fotograma.
                with state.lock:
                    if (
                        state.detection_active
                        and not state.inference_busy
                        and state.inference_generation == inference_generation
                    ):
                        state.inference_busy = True
                        inference_busy = True
                        Thread(
                            target=run_live_inference,
                            args=(image.copy(), inference_generation),
                            daemon=True,
                        ).start()

            # En la cámara en vivo siempre se muestran las dos clases. Los
            # filtros de máscaras pertenecen únicamente al análisis de fotos.
            mask_layers = [
                cached_mask
                for cached_mask in (cached_healthy_masks, cached_sick_masks)
                if cached_mask is not None
            ]
            if not mask_layers:
                output = image
            elif len(mask_layers) == 1:
                output = overlay_masks(image, mask_layers[0], mask_opacity)
            else:
                output = overlay_masks(
                    image,
                    cv2.bitwise_or(mask_layers[0], mask_layers[1]),
                    mask_opacity,
                )
            output = draw_live_status_overlay(
                output,
                True,
                copy_output=False,
            )
            if frame_number % LIVE_RECORD_EVERY_N_FRAMES == 1:
                state.record_frame(
                    output,
                    resolution["frame_rate"] / LIVE_RECORD_EVERY_N_FRAMES,
                )
            return make_output_frame(output)
        except Exception as error:
            with state.lock:
                state.last_error = f"{type(error).__name__}: {error}"
            # Si un cuadro puntual falla, se conserva el video en lugar de cerrar la cámara.
            if image is not None:
                return make_output_frame(image)
            return frame

    with camera_action:
        st.button(
            "Detener cámara" if camera_requested else "Iniciar cámara",
            type="secondary" if camera_requested else "primary",
            key="live_camera_toggle",
            use_container_width=True,
            on_click=toggle_live_camera,
        )

    camera_column, status_column = st.columns([1.55, 1.15], gap="large")
    with camera_column:
        # WebRTC se encarga de transportar y pintar el video a la frecuencia
        # de la cámara. Streamlit no debe reconstruir una imagen JPEG en cada
        # rerun: hacerlo por el WebSocket de la interfaz causa saltos visibles.
        camera_playing = False
        ice_state = ""
        if camera_requested:
            webrtc_options: dict[str, Any] = {
                "key": "cell_live_camera",
                "mode": WebRtcMode.SENDRECV,
                "desired_playing_state": True,
                "rtc_configuration": rtc_configuration(),
                "media_stream_constraints": {
                    "video": {
                        "width": {"ideal": resolution["width"]},
                        "height": {"ideal": resolution["height"]},
                        "frameRate": {"ideal": resolution["frame_rate"]},
                    },
                    "audio": False,
                },
                "video_frame_callback": process_live_frame,
                # El callback solo entrega el último cuadro y agenda YOLO en
                # otro hilo. La cola de un elemento evita que la inferencia
                # lenta acumule retraso y memoria mientras la cámara continúa.
                "async_processing": True,
                "video_receiver_size": 1,
                "sendback_video": True,
                "sendback_audio": False,
            }
            camera_context = webrtc_streamer(**webrtc_options)
            camera_state = getattr(camera_context, "state", None)
            camera_playing = bool(getattr(camera_state, "playing", False))
            ice_state = str(getattr(camera_state, "ice_connection_state", ""))
            st.session_state["live_camera_playing"] = camera_playing
            if camera_playing:
                st.session_state.pop("live_camera_start_required", None)
        else:
            st.session_state["live_camera_playing"] = False

        camera_snapshot = state.snapshot()
        with status_column:
            render_live_metrics_panel(
                state,
                lot_name,
                sample_code,
                10.0,
                30.0,
                60.0,
            )

    detection_snapshot = state.snapshot()
    detection_is_active = detection_snapshot["detection_active"] or bool(
        st.session_state.get("live_detection_requested", False)
    )
    sample_waiting_to_save = bool(
        detection_snapshot["sample_ready"] and not detection_is_active
    )
    with detection_action:
        detection_label = "Detener detección" if detection_is_active else "Iniciar detección"
        st.button(
            detection_label,
            type="primary",
            key="live_detection_toggle",
            use_container_width=True,
            disabled=(
                completed_count >= MAX_LIVE_SAMPLES
                or not camera_requested
                or sample_waiting_to_save
            ),
            on_click=toggle_live_detection,
        )
    if state.snapshot()["detection_active"] or st.session_state.get(
        "live_detection_requested", False
    ):
        ensure_live_model()
    with save_action:
        snapshot_after_action = state.snapshot()
        if st.button(
            "Guardar muestra",
            key="live_sample_save",
            use_container_width=True,
            disabled=(
                detection_is_active
                or completed_count >= MAX_LIVE_SAMPLES
                or not snapshot_after_action["sample_ready"]
            ),
        ):
            state.set_detection_active(False)
            st.session_state["live_detection_requested"] = False
            if lot_name:
                remember_pool(lot_name)
            if state.save_current_sample(sample_code, lot_name):
                persist_completed_live_samples(state)
                saved_samples = state.completed_samples_snapshot()
                if saved_samples:
                    add_live_sample_to_history(saved_samples[-1])
                st.session_state.pop("live_pending_sample", None)
                st.session_state["live_save_feedback"] = "Muestra guardada correctamente."
                st.rerun()
            else:
                st.warning("No hay una captura detenida lista para guardar.")

        if snapshot_after_action["sample_ready"] and not detection_is_active:
            if st.button(
                "Borrar muestra",
                key="live_sample_discard",
                use_container_width=True,
            ):
                state.set_detection_active(False)
                state.discard_pending_sample()
                st.session_state["live_detection_requested"] = False
                st.session_state.pop("live_pending_sample", None)
                st.session_state["live_save_feedback"] = (
                    "Muestra descartada. Puedes iniciar una nueva detección."
                )
                st.rerun()

    render_live_sample_results(
        state,
        lot_name,
        camera_playing,
        10.0,
        30.0,
        60.0,
    )


def main() -> None:
    st.set_page_config(page_title="Análisis celular", layout="wide")
    st.markdown(
        """
        <style>
            .stApp { background: #061a33; color: #ffffff; }
            [data-testid="stHeader"] { background: rgba(0, 0, 0, 0); }
            [data-testid="stSidebar"] { background: #0a2745; }
            /* El tema base reserva demasiado espacio antes y después del
               contenido. Mantener una separación segura del encabezado, pero
               evitar que la página termine con un vacío artificial. */
            [data-testid="stMainBlockContainer"] {
                padding-top: 4.5rem !important;
                padding-bottom: 3rem !important;
            }
            h1, h2, h3, h4, p, label, .stMarkdown, .stCaption,
            [data-testid="stMetricLabel"], [data-testid="stMetricValue"],
            [data-testid="stMetricDelta"] { color: #ffffff !important; }
            [data-testid="stMetric"] {
                background: #103858;
                border: 1px solid #41d8cc;
                border-radius: 10px;
                padding: 0.75rem;
            }
            [data-testid="stFileUploader"] {
                background: #0d3153;
                border: 1px solid #41d8cc;
                border-radius: 10px;
                padding: 0.2rem 0.35rem;
            }
            [data-testid="stFileUploaderDropzone"] {
                min-height: 3.1rem !important;
                padding: 0.25rem 0.55rem !important;
                background: #0d3153 !important;
                border: 1px dashed #6be7da !important;
                border-radius: 7px !important;
            }
            /* Una vez cargada la foto, solo queda su nombre en una fila compacta. */
            [data-testid="stFileUploader"]:has([data-testid="stFileUploaderFile"])
            [data-testid="stFileUploaderDropzone"] {
                display: none !important;
            }
            [data-testid="stFileUploaderFile"] {
                min-height: 2.2rem !important;
                margin: 0 !important;
                padding: 0.25rem 0.4rem !important;
                font-size: 0.85rem !important;
            }
            [data-testid="stFileUploaderFile"] button {
                min-height: 1.8rem !important;
                min-width: 1.8rem !important;
                padding: 0.15rem !important;
            }
            [data-testid="stFileUploaderDropzone"] svg {
                width: 1.45rem !important;
                height: 1.45rem !important;
                color: #41d8cc !important;
            }
            [data-testid="stFileUploaderDropzoneInstructions"] {
                margin: 0 !important;
            }
            [data-testid="stFileUploaderDropzoneInstructions"] div {
                font-size: 0.75rem !important;
                color: #ffffff !important;
                opacity: 1 !important;
            }
            [data-testid="stFileUploaderDropzoneInstructions"] *,
            [data-testid="stFileUploaderFile"], [data-testid="stFileUploaderFile"] * {
                color: #ffffff !important;
                opacity: 1 !important;
            }
            [data-testid="stFileUploaderDropzone"] button {
                background: #12b8c2 !important;
                color: #ffffff !important;
                border: none !important;
                font-weight: 700 !important;
                min-height: 2rem !important;
                padding: 0.25rem 0.75rem !important;
            }
            [data-testid="stFileUploaderDropzone"] button * {
                color: #ffffff !important;
            }
            .upload-heading {
                display: flex;
                align-items: center;
                gap: 0.65rem;
                margin: 0.15rem 0 0.45rem;
                color: #ffffff;
            }
            .upload-heading strong { display: block; font-size: 1rem; }
            .upload-heading small { display: block; color: #a6d8eb; font-size: 0.78rem; }
            .upload-step {
                display: inline-flex;
                align-items: center;
                justify-content: center;
                width: 1.65rem;
                height: 1.65rem;
                border-radius: 50%;
                background: #12b8c2;
                color: #ffffff;
                font-weight: 800;
            }
            .photo-file-row {
                box-sizing: border-box;
                min-height: 2.75rem;
                display: flex;
                align-items: center;
                gap: 0.5rem;
                margin: 0 !important;
                padding: 0.45rem 0.75rem;
                background: #0d3153;
                border: 1px solid #41d8cc;
                border-radius: 8px;
                color: #ffffff;
                overflow: hidden;
            }
            .photo-file-row span {
                flex: 0 0 auto;
                color: #a6d8eb;
                font-size: 0.78rem;
            }
            .photo-file-row strong {
                overflow: hidden;
                text-overflow: ellipsis;
                white-space: nowrap;
                font-size: 0.9rem;
            }
            [data-testid="stAlert"] { color: #ffffff; }
            .stButton > button, .stDownloadButton > button,
            button[kind="primary"], [data-testid="stBaseButton-primary"] {
                background: #0f7c86 !important;
                background-color: #0f7c86 !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                border: 1px solid #8ef4e8 !important;
                font-weight: 750 !important;
                min-height: 2.75rem;
            }
            .stButton > button:hover, .stDownloadButton > button:hover,
            button[kind="primary"]:hover, [data-testid="stBaseButton-primary"]:hover,
            .stButton > button:focus-visible, .stDownloadButton > button:focus-visible,
            button[kind="primary"]:focus-visible, [data-testid="stBaseButton-primary"]:focus-visible {
                background: #176b8a !important;
                background-color: #176b8a !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                border-color: #c4fffa !important;
            }
            .stButton > button:active, .stDownloadButton > button:active,
            button[kind="primary"]:active, [data-testid="stBaseButton-primary"]:active,
            button[kind="secondary"]:active, [data-testid="stBaseButton-secondary"]:active,
            button[kind="secondaryFormSubmit"]:active,
            [data-testid="stBaseButton-secondaryFormSubmit"]:active {
                background: #176b8a !important;
                background-color: #176b8a !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                border-color: #c4fffa !important;
            }
            button[kind="secondary"], [data-testid="stBaseButton-secondary"],
            button[kind="secondaryFormSubmit"],
            [data-testid="stBaseButton-secondaryFormSubmit"] {
                background: #31546b !important;
                background-color: #31546b !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                border: 1px solid #6bbfd1 !important;
                font-weight: 750 !important;
                min-height: 2.75rem;
            }
            button[kind="secondary"]:hover, [data-testid="stBaseButton-secondary"]:hover,
            button[kind="secondary"]:focus-visible, [data-testid="stBaseButton-secondary"]:focus-visible,
            button[kind="secondaryFormSubmit"]:hover,
            [data-testid="stBaseButton-secondaryFormSubmit"]:hover,
            button[kind="secondaryFormSubmit"]:focus-visible,
            [data-testid="stBaseButton-secondaryFormSubmit"]:focus-visible {
                background: #176b8a !important;
                background-color: #176b8a !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                border-color: #c4fffa !important;
            }
            .stButton > button *, .stDownloadButton > button *,
            button[kind="primary"] *, [data-testid="stBaseButton-primary"] * {
                color: #ffffff !important;
            }
            button[kind="secondary"] *, [data-testid="stBaseButton-secondary"] *,
            button[kind="secondaryFormSubmit"] *,
            [data-testid="stBaseButton-secondaryFormSubmit"] * {
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            .stButton > button:disabled, button[kind="primary"]:disabled,
            button[kind="secondary"]:disabled,
            [data-testid="stBaseButton-secondary"]:disabled,
            button[kind="secondaryFormSubmit"]:disabled,
            [data-testid="stBaseButton-secondaryFormSubmit"]:disabled {
                background: #31546b !important;
                background-color: #31546b !important;
                border-color: #31546b !important;
                color: #b6c8d5 !important;
                -webkit-text-fill-color: #b6c8d5 !important;
                opacity: 1 !important;
            }
            input, textarea { color: #ffffff !important; background: #0d3153 !important; }
            /* Streamlit/BaseWeb vuelve a pintar estos controles con fondo blanco
               durante los reruns de la cámara. Mantener contraste en todo momento. */
            [data-testid="stTextInput"] > div > div,
            [data-testid="stTextInput"] input {
                background: #0d3153 !important;
                color: #ffffff !important;
                border-color: #6bbfd1 !important;
                caret-color: #ffffff !important;
            }
            [data-testid="stTextInput"] [data-baseweb="input"],
            [data-testid="stTextInput"] [data-baseweb="base-input"],
            [data-testid="stTextInput"] [data-baseweb="input"] > div,
            [data-testid="stTextInput"] [data-baseweb="base-input"] > div,
            [data-testid="stSelectbox"] [data-baseweb="select"],
            [data-testid="stSelectbox"] [data-baseweb="select"] > div {
                background-color: #0d3153 !important;
                border-color: #6bbfd1 !important;
            }
            [data-testid="stTextInput"] input::placeholder {
                color: #a6d8eb !important;
                opacity: 1 !important;
            }
            [data-baseweb="input"],
            [data-baseweb="input"] > div,
            [data-baseweb="base-input"],
            [data-baseweb="base-input"] > div {
                background: #0d3153 !important;
                border-color: #6bbfd1 !important;
            }
            [data-baseweb="input"] input,
            [data-baseweb="base-input"] input {
                background: transparent !important;
                color: #ffffff !important;
                caret-color: #ffffff !important;
            }
            [data-testid="stSelectbox"] [data-baseweb="select"] > div {
                background: #0d3153 !important;
                border-color: #6bbfd1 !important;
                color: #ffffff !important;
            }
            [data-testid="stSelectbox"] [data-baseweb="select"] > div > div {
                background: #0d3153 !important;
            }
            [data-testid="stSelectbox"] [data-baseweb="select"] div,
            [data-testid="stSelectbox"] [data-baseweb="select"] span,
            [data-testid="stSelectbox"] [data-baseweb="select"] input {
                color: #ffffff !important;
            }
            [data-testid="stSelectbox"] [data-baseweb="select"] svg {
                fill: #c4fffa !important;
            }
            /* Pestañas de navegación: el estado activo y el hover conservan
               el mismo contraste que el resto de la interfaz. */
            [data-testid="stTabs"] [role="tablist"] {
                gap: 0.35rem;
                border-bottom: 1px solid #2c8396;
            }
            [data-testid="stTabs"] [role="tab"] {
                background: #0d3153 !important;
                border: 1px solid #2c8396 !important;
                border-bottom: 3px solid transparent !important;
                border-radius: 8px 8px 0 0 !important;
                color: #c8eaf4 !important;
                -webkit-text-fill-color: #c8eaf4 !important;
                font-weight: 750 !important;
                padding: 0.55rem 1rem !important;
            }
            [data-testid="stTabs"] [role="tab"]:hover,
            [data-testid="stTabs"] [role="tab"]:focus-visible,
            [data-testid="stTabs"] [role="tab"][aria-selected="true"] {
                background: #176b8a !important;
                background-color: #176b8a !important;
                border-color: #c4fffa !important;
                border-bottom-color: #c4fffa !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                outline: none !important;
            }
            [data-testid="stTabs"] [role="tab"] * {
                color: inherit !important;
                -webkit-text-fill-color: inherit !important;
            }
            [data-testid="stTabs"] [data-baseweb="tab-panel"] {
                padding-top: 0.85rem !important;
            }
            /* Expanders usan una regla BaseWeb propia que, sin esta
               sobreescritura, vuelve a pintar el encabezado de blanco. */
            [data-testid="stExpander"] {
                background: #0a2745 !important;
                border: 1px solid #2c8396 !important;
                border-radius: 10px !important;
            }
            [data-testid="stExpander"] details,
            [data-testid="stExpander"] summary {
                background: #0d3153 !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            [data-testid="stExpander"] summary {
                border-radius: 9px !important;
                padding: 0.65rem 0.8rem !important;
            }
            [data-testid="stExpander"] summary:hover,
            [data-testid="stExpander"] details[open] summary {
                background: #176b8a !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            [data-testid="stExpander"] summary * {
                color: inherit !important;
                -webkit-text-fill-color: inherit !important;
            }
            [data-testid="stExpander"] summary svg {
                fill: #c4fffa !important;
                color: #c4fffa !important;
            }
            [data-testid="stExpander"] details > div {
                background: #061a33 !important;
                color: #ffffff !important;
            }
            [data-baseweb="popover"] [role="listbox"],
            [data-baseweb="popover"] [role="option"] {
                background: #0d3153 !important;
                color: #ffffff !important;
            }
            [data-baseweb="popover"] [role="option"]:hover,
            [data-baseweb="popover"] [aria-selected="true"] {
                background: #176b8a !important;
                color: #ffffff !important;
            }
            [data-testid="stRadio"] div[role="radiogroup"] { gap: 0.5rem; }
            [data-testid="stRadio"] label {
                background: #0d3153;
                border: 1px solid #2a7895;
                border-radius: 8px;
                color: #ffffff !important;
                padding: 0.4rem 1rem;
            }
            [data-testid="stRadio"] label:has(input:checked) {
                background: #176b8a;
                border-color: #6bbfd1;
                color: #ffffff !important;
            }
            [data-testid="stRadio"] label:has(input:checked) p { color: #ffffff !important; }
            [data-testid="stRadio"] label:has(input:checked) > div:first-child,
            [data-testid="stToggle"] label:has(input:checked) > div:first-child {
                background: #12b8c2 !important;
                border-color: #12b8c2 !important;
            }
            [data-testid="stRadio"] label:has(input:checked) > div:first-child > div,
            [data-testid="stToggle"] label:has(input:checked) > div:first-child > div {
                background: #ffffff !important;
            }
            [data-testid="stRadio"] label p, [data-testid="stToggle"] label,
            [data-testid="stToggle"] p { color: #ffffff !important; }
            button[data-testid="stBaseButton-pills"],
            button[data-testid="stBaseButton-pillsActive"],
            [data-testid="stPills"] button {
                background: #0d3153 !important;
                border: 1px solid #41d8cc !important;
                color: #ffffff !important;
                border-radius: 999px !important;
            }
            button[data-testid="stBaseButton-pillsActive"],
            [data-testid="stPills"] button[aria-pressed="true"] {
                background: #176b8a !important;
                background-color: #176b8a !important;
                border-color: #6bbfd1 !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            [data-testid="stPills"] button * { color: inherit !important; }
            /* En Streamlit 1.64 los pills se renderizan como stButtonGroup.
               Cubrir el contenedor real evita el rojo/blanco del tema base. */
            [data-testid="stButtonGroup"] button {
                background: #0d3153 !important;
                background-color: #0d3153 !important;
                border: 1px solid #41d8cc !important;
                border-radius: 999px !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                box-shadow: none !important;
            }
            [data-testid="stButtonGroup"] button:hover,
            [data-testid="stButtonGroup"] button:focus-visible {
                background: #176b8a !important;
                background-color: #176b8a !important;
                border-color: #c4fffa !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                outline: none !important;
            }
            [data-testid="stButtonGroup"] button[aria-pressed="true"],
            [data-testid="stButtonGroup"] button[aria-pressed="true"]:hover,
            [data-testid="stButtonGroup"] button[aria-pressed="true"]:focus-visible {
                background: #176b8a !important;
                background-color: #176b8a !important;
                border-color: #c4fffa !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            [data-testid="stButtonGroup"] button * { color: inherit !important; }
            /* El video debe permanecer en WebRTC: así conserva la frecuencia
               de la cámara y no depende de reruns del WebSocket de Streamlit. */
            [data-testid="stCustomComponentV1"] {
                width: min(100%, 27rem) !important;
                max-width: 27rem !important;
                min-height: 24rem !important;
                height: 24rem !important;
                overflow: visible !important;
                margin: 0.75rem auto 0 !important;
                padding: 0 !important;
            }
            [data-testid="stCustomComponentV1"] iframe {
                display: block !important;
                position: relative !important;
                width: 100% !important;
                min-height: 24rem !important;
                max-height: 24rem !important;
                height: 24rem !important;
                border: 0 !important;
            }
            .live-preview {
                max-width: 42rem;
                margin: 0.25rem 0 0.8rem;
                padding: 0.55rem 0.7rem;
                background: #0a2745;
                border: 1px solid #41d8cc;
                border-radius: 10px;
            }
            .live-preview-header {
                display: flex;
                align-items: center;
                gap: 0.45rem;
                color: #ffffff;
                font-size: 0.9rem;
            }
            .live-preview-dot {
                width: 0.55rem;
                height: 0.55rem;
                border-radius: 50%;
                background: #12b8c2;
                box-shadow: 0 0 0 3px rgba(18, 184, 194, 0.18);
            }
            .live-preview-state {
                margin-left: auto;
                color: #a6f5f0;
                font-size: 0.72rem;
                font-weight: 700;
            }
            .live-preview-screen {
                display: flex;
                align-items: center;
                justify-content: center;
                min-height: 3.8rem;
                margin-top: 0.5rem;
                padding: 0.45rem;
                background: #061a33;
                border: 1px dashed #2c8396;
                border-radius: 7px;
                color: #a6d8eb;
                font-size: 0.82rem;
                text-align: center;
            }
            .live-video-placeholder {
                display: flex;
                align-items: center;
                justify-content: center;
                min-height: 24rem;
                padding: 1rem;
                background: #061a33;
                border: 1px solid #2c8396;
                border-radius: 12px;
                color: #a6d8eb;
                text-align: center;
            }
            .live-preview-footer {
                margin-top: 0.35rem;
                color: #a6d8eb;
                font-size: 0.72rem;
            }
            .live-measurements-heading {
                margin: 1rem 0 0.45rem;
                color: #a6f5f0;
                font-size: 0.88rem;
                font-weight: 800;
                text-transform: uppercase;
                letter-spacing: 0.04em;
            }
            .live-measurement-list {
                display: grid;
                gap: 0.55rem;
                margin-bottom: 0.75rem;
            }
            .live-measurement-list div {
                display: flex;
                flex-direction: column;
                gap: 0.15rem;
                padding: 0.55rem 0.7rem;
                background: #0d3153;
                border: 1px solid rgba(107, 191, 209, 0.58);
                border-radius: 8px;
            }
            .live-measurement-list strong { color: #ffffff; font-size: 0.82rem; }
            .live-measurement-list span { color: #e8ffff; font-size: 0.78rem; line-height: 1.35; }
            .live-sample-cards {
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                max-width: 27rem;
                gap: 0.45rem;
                margin-top: 0.7rem;
                margin-left: auto;
            }
            .live-sample-card {
                min-width: 0;
                padding: 0.6rem;
                background: #0d3153;
                border: 1px solid rgba(107, 191, 209, 0.62);
                border-radius: 10px;
                color: #ffffff;
            }
            .live-current-card {
                border-color: #c4fffa;
                box-shadow: 0 0 0 1px rgba(196, 255, 250, 0.18);
            }
            .live-aggregate-card {
                background: #176b8a;
                border: 2px solid #c4fffa;
            }
            .live-sample-card-heading {
                display: flex;
                flex-direction: column;
                gap: 0.15rem;
                margin-bottom: 0.55rem;
            }
            .live-sample-card-heading strong {
                color: #ffffff;
                font-size: 0.86rem;
                line-height: 1.2;
            }
            .live-sample-card-heading span {
                color: #a6f5f0;
                font-size: 0.7rem;
                overflow-wrap: anywhere;
            }
            .live-card-grid {
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 0.3rem;
            }
            .live-card-grid div {
                display: flex;
                flex-direction: column;
                gap: 0.1rem;
                padding: 0.35rem;
                background: rgba(6, 26, 51, 0.35);
                border-radius: 6px;
            }
            .live-card-grid span, .live-card-grade span {
                color: #a6d8eb;
                font-size: 0.66rem;
            }
            .live-card-grid strong {
                color: #ffffff;
                font-size: 0.9rem;
            }
            .live-card-grade {
                display: flex;
                align-items: center;
                justify-content: space-between;
                gap: 0.4rem;
                margin-top: 0.45rem;
                padding-top: 0.45rem;
                border-top: 1px solid rgba(107, 231, 218, 0.22);
            }
            .live-card-grade strong {
                color: #c4fffa;
                font-size: 0.8rem;
                text-align: right;
            }
            .live-card-measurements {
                margin-top: 0.35rem;
                color: #e8ffff;
                font-size: 0.61rem;
                line-height: 1.35;
            }
            .live-empty-metrics {
                margin: 0.7rem 0 0;
                color: #a6d8eb;
                font-size: 0.82rem;
            }
            .live-overall-summary {
                display: flex;
                flex-direction: column;
                gap: 0.25rem;
                margin: 0.85rem 0;
                padding: 0.9rem 1rem;
                background: #176b8a;
                border: 2px solid #c4fffa;
                border-radius: 12px;
                color: #ffffff;
            }
            .live-overall-summary strong { font-size: 1.05rem; }
            .live-overall-summary span { color: #e8ffff; font-size: 0.88rem; }
            @media (max-width: 1100px) {
                .live-sample-cards { grid-template-columns: 1fr; }
            }
            .mask-legend {
                display: inline-flex;
                align-items: center;
                flex-wrap: wrap;
                gap: 0.45rem 1.1rem;
                margin: 0.5rem 0 0.75rem;
                padding: 0.48rem 0.85rem;
                background: linear-gradient(135deg, rgba(23, 107, 138, 0.58), rgba(10, 39, 69, 0.78));
                backdrop-filter: blur(12px);
                -webkit-backdrop-filter: blur(12px);
                border: 1px solid rgba(107, 191, 209, 0.62);
                border-radius: 9px;
                color: #ffffff;
                font-size: 0.85rem;
                box-shadow: 0 4px 12px rgba(0, 0, 0, 0.16);
            }
            .mask-legend span {
                display: inline-flex;
                align-items: center;
                gap: 0.35rem;
                color: #ffffff;
                font-weight: 650;
            }
            .legend-dot {
                display: inline-block;
                width: 0.65rem;
                height: 0.65rem;
                border-radius: 50%;
            }
            .legend-healthy { background: #48b848; }
            .legend-sick { background: #eb3c3c; }
            .results-panel {
                background: #0a2745;
                border: 1px solid #41d8cc;
                border-radius: 16px;
                padding: 1rem;
                margin-top: 0.5rem;
            }
            .results-heading {
                color: #ffffff;
                font-size: 1.7rem;
                font-weight: 750;
                margin-bottom: 1.1rem;
            }
            .results-grid {
                display: grid;
                grid-template-columns: 1fr;
                gap: 1rem;
            }
            .result-card {
                min-width: 0;
                background: #176b8a;
                border: 1px solid #6bbfd1;
                border-radius: 12px;
                min-height: 4.9rem;
                padding: 1.1rem 1.15rem;
                display: flex;
                align-items: center;
                justify-content: space-between;
                gap: 0.75rem;
            }
            .result-label {
                display: block;
                color: #ffffff;
                font-size: 1.05rem;
                font-weight: 650;
                line-height: 1.2;
                min-height: 0;
                overflow-wrap: anywhere;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .result-value {
                display: block;
                color: #ffffff;
                font-size: 2.5rem;
                line-height: 1;
                font-weight: 700;
                text-align: right;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .measurement-card {
                min-height: 0;
                display: block;
                padding: 1rem 1.15rem;
            }
            .measurement-heading {
                display: block;
                color: #ffffff;
                font-size: 1.05rem;
                font-weight: 750;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .measurement-row {
                display: flex;
                align-items: baseline;
                justify-content: space-between;
                gap: 0.75rem;
                margin-top: 0.5rem;
                padding-top: 0.5rem;
                border-top: 1px solid rgba(255, 255, 255, 0.45);
            }
            .measurement-row span {
                color: #e8ffff;
                font-size: 0.82rem;
                font-weight: 650;
            }
            .measurement-row strong {
                color: #ffffff;
                font-size: 1.05rem;
                font-weight: 800;
                text-align: right;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .measurement-empty {
                margin-top: 0.5rem;
                color: #e8ffff;
                font-size: 0.9rem;
            }
            .measurement-inline {
                display: flex;
                align-items: center;
                flex-wrap: wrap;
                gap: 0.4rem 1rem;
                margin: 0.7rem 0 0.25rem;
                padding: 0.65rem 0.85rem;
                background: #0d3153;
                border: 1px solid #41d8cc;
                border-radius: 9px;
                color: #ffffff;
            }
            .measurement-inline-title {
                color: #a6f5f0;
                font-size: 0.85rem;
                font-weight: 700;
            }
            .measurement-inline strong {
                color: #ffffff;
                font-size: 1rem;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .grade-card { background: #1b7894; flex-wrap: wrap; }
            .grade-featured {
                min-height: 8.1rem;
                padding: 1.3rem 1.25rem;
                border: 2px solid #c4fffa;
                box-shadow: 0 7px 18px rgba(0, 0, 0, 0.18);
            }
            .grade-featured .result-label { font-size: 1.2rem; }
            .grade-featured .grade-value { font-size: 2.45rem; }
            .grade-featured .grade-detail { font-size: 1.05rem; margin-top: 0.15rem; }
            .grade-value {
                display: block;
                color: #ffffff;
                font-size: 1.45rem;
                font-weight: 750;
                line-height: 1.2;
                margin-left: auto;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .grade-detail {
                display: inline-block;
                color: #ffffff;
                font-size: 0.9rem;
                font-weight: 600;
                flex-basis: 100%;
                margin-top: -0.25rem;
                text-shadow: 0 2px 4px rgba(6, 26, 51, 0.62);
            }
            .results-note {
                color: #e7f2ff !important;
                font-size: 0.9rem;
                line-height: 1.45;
                margin: 3.25rem 0 0 !important;
                padding-top: 0.25rem;
            }
            /* En móvil Streamlit apila las columnas, pero conserva el gap de
               4rem usado en escritorio. Reducir solo el espacio vertical
               evita huecos entre los campos y entre el video y sus métricas. */
            @media (max-width: 768px) {
                [data-testid="stHorizontalBlock"] {
                    row-gap: 0.75rem !important;
                }
                [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] {
                    min-width: 0 !important;
                }
            }

            /* Capa visual de aplicación: jerarquía clara, paneles compactos
               y navegación que se adapta a escritorio y móvil. */
            [data-testid="stMainBlockContainer"] {
                width: min(100%, 92rem) !important;
                max-width: 92rem !important;
                margin: 0 auto !important;
            }
            .app-header {
                position: relative;
                display: grid;
                grid-template-columns: 1fr;
                align-items: stretch;
                margin: 0 0 0.9rem;
                padding: 0.95rem 0.25rem 1.25rem;
                overflow: visible;
                background: transparent;
                border: 0;
                border-bottom: 1px solid rgba(107, 191, 209, 0.3);
                border-radius: 0;
                box-shadow: none;
            }
            .app-header-copy { position: relative; z-index: 1; }
            .app-eyebrow {
                margin-bottom: 0.38rem;
                color: #a6f5f0;
                font-size: 0.7rem;
                font-weight: 800;
                letter-spacing: 0.14em;
                text-transform: uppercase;
            }
            .app-title {
                margin: 0 !important;
                color: #ffffff !important;
                font-size: clamp(2rem, 4.2vw, 3.45rem) !important;
                font-weight: 800 !important;
                letter-spacing: -0.045em;
                line-height: 1.04 !important;
            }
            .app-title-accent { color: #8ef4e8 !important; }
            .app-subtitle {
                max-width: 52rem;
                margin: 0.55rem 0 0 !important;
                color: #c8eaf4 !important;
                font-size: clamp(0.82rem, 1.4vw, 0.98rem) !important;
                line-height: 1.45 !important;
            }
            .history-count {
                margin: 0.65rem 0 0.75rem;
                color: #a6f5f0;
                font-size: 0.78rem;
                font-weight: 750;
                letter-spacing: 0.04em;
                text-transform: uppercase;
            }
            .history-empty {
                display: flex;
                flex-direction: column;
                gap: 0.3rem;
                margin-top: 0.9rem;
                padding: 1.1rem 1.2rem;
                background: #0a2745;
                border: 1px solid #2c8396;
                border-radius: 12px;
                color: #ffffff;
            }
            .history-empty span { color: #a6d8eb; font-size: 0.86rem; }
            [class*="st-key-history_row_"] {
                margin-top: 0.75rem;
                padding: 0.3rem 0 0.55rem;
            }
            [class*="st-key-history_row_"] [data-testid="stHorizontalBlock"] {
                align-items: center;
            }
            [class*="st-key-history_row_"] [data-testid="stColumn"]:first-child {
                max-width: 250px;
            }
            [class*="st-key-history_preview_"] {
                position: relative !important;
                overflow: hidden;
                border: 1px solid rgba(107, 191, 209, 0.5);
                border-radius: 10px;
                background: #0a2745;
            }
            [class*="st-key-history_preview_"] [data-testid="stImage"] {
                margin: 0 !important;
            }
            [class*="st-key-history_preview_"] [class*="st-key-history_play_"] {
                position: absolute !important;
                inset: 0 !important;
                z-index: 3;
                display: flex !important;
                align-items: center;
                justify-content: center;
                width: 100% !important;
                height: 100% !important;
            }
            [class*="st-key-history_preview_"] [class*="st-key-history_play_"] button {
                width: 2.65rem !important;
                min-width: 2.65rem !important;
                height: 2.65rem !important;
                min-height: 2.65rem !important;
                padding: 0 !important;
                border: 1px solid rgba(177, 255, 247, 0.9) !important;
                border-radius: 50% !important;
                background: rgba(6, 38, 68, 0.88) !important;
                color: #ffffff !important;
                font-size: 1.1rem !important;
                box-shadow: 0 3px 14px rgba(0, 0, 0, 0.35);
                pointer-events: auto;
            }
            [class*="st-key-history_preview_"] [data-testid="stVideo"] {
                margin: 0 !important;
                overflow: hidden;
                border-radius: 9px;
            }
            .history-video-missing {
                display: grid;
                aspect-ratio: 16 / 9;
                place-items: center;
                padding: 0.35rem;
                background: #0a2745;
                color: #a6d8eb;
                font-size: 0.68rem;
                text-align: center;
            }
            .history-row-meta {
                min-width: 0;
                padding: 0.4rem 0.2rem;
            }
            .history-row-top {
                display: flex;
                align-items: center;
                justify-content: space-between;
                gap: 0.65rem;
            }
            .history-row-code {
                min-width: 0;
                overflow-wrap: anywhere;
                color: #ffffff;
                font-size: clamp(0.95rem, 2.5vw, 1.18rem);
                font-weight: 750;
            }
            .history-row-details {
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 0.7rem;
                margin-top: 0.65rem;
            }
            .history-row-details span { min-width: 0; }
            .history-row-details small,
            .history-row-details strong {
                display: block;
                overflow-wrap: anywhere;
            }
            .history-row-details small {
                color: #a6d8eb;
                font-size: 0.68rem;
            }
            .history-row-details strong {
                margin-top: 0.1rem;
                color: #ffffff;
                font-size: 0.82rem;
                font-weight: 600;
            }
            .history-row-divider {
                height: 1px;
                margin: 0.1rem 0 0.35rem;
                background: rgba(107, 191, 209, 0.27);
            }
            .history-card-grade {
                flex: 0 0 auto;
                padding: 0.28rem 0.55rem;
                background: #176b8a;
                border: 1px solid #6bbfd1;
                border-radius: 999px;
                color: #ffffff;
                font-size: 0.75rem;
                font-weight: 800;
            }
            @media (max-width: 380px) {
                [class*="st-key-history_row_"] [data-testid="stHorizontalBlock"] {
                    flex-direction: column !important;
                    align-items: stretch !important;
                }
                [class*="st-key-history_row_"] [data-testid="stColumn"] {
                    width: 100% !important;
                    max-width: 100% !important;
                    flex: 1 1 100% !important;
                }
                [class*="st-key-history_preview_"] { max-width: 250px; }
            }
            [data-testid="stTabs"] {
                margin: 0.35rem 0 1.1rem !important;
            }
            [data-testid="stTabs"] [role="tablist"] {
                display: flex !important;
                flex-wrap: wrap !important;
                gap: 0.35rem !important;
                padding: 0.35rem !important;
                overflow: visible !important;
                background: rgba(10, 39, 69, 0.82) !important;
                border: 1px solid rgba(107, 191, 209, 0.4) !important;
                border-radius: 13px !important;
                box-shadow: 0 8px 22px rgba(0, 0, 0, 0.12);
            }
            [data-testid="stTabs"] [role="tab"] {
                flex: 0 1 auto !important;
                min-height: 2.45rem !important;
                padding: 0.58rem 0.95rem !important;
                background: transparent !important;
                border: 1px solid transparent !important;
                border-radius: 9px !important;
                color: #c8eaf4 !important;
                -webkit-text-fill-color: #c8eaf4 !important;
                font-size: 0.88rem !important;
                font-weight: 750 !important;
                white-space: nowrap !important;
                transition: background 160ms ease, border-color 160ms ease,
                    color 160ms ease, transform 160ms ease !important;
            }
            [data-testid="stTabs"] [role="tab"]:hover,
            [data-testid="stTabs"] [role="tab"]:focus-visible,
            [data-testid="stTabs"] [role="tab"][aria-selected="true"] {
                background: linear-gradient(135deg, #176b8a, #0f7c86) !important;
                background-color: #176b8a !important;
                border-color: #8ef4e8 !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
                outline: none !important;
            }
            [data-testid="stTabs"] [role="tab"][aria-selected="true"] {
                box-shadow: 0 4px 12px rgba(15, 124, 134, 0.28);
            }
            [data-testid="stTabs"] [role="tab"]:hover {
                transform: translateY(-1px);
            }
            [data-testid="stTabs"] [role="tab"]::after {
                display: none !important;
            }
            [data-testid="stTabs"] .react-aria-SelectionIndicator {
                background: #c4fffa !important;
                border: 0 !important;
                border-radius: 999px !important;
                height: 0.2rem !important;
            }
            [data-testid="stTabs"] [role="tab"] * {
                color: inherit !important;
                -webkit-text-fill-color: inherit !important;
            }
            /* Las pestañas usan el estado activo como pastilla; no dibujar
               una línea inferior adicional. */
            [data-testid="stTabs"] [role="tab"] {
                border-bottom: 0 !important;
            }
            [data-testid="stTabs"] .react-aria-SelectionIndicator {
                display: none !important;
            }
            /* Navegación principal: una fila de pastillas oscura, sin paneles
               blancos, coherente con la paleta de la aplicación. Streamlit la
               renderiza como un grupo de radios cuando se usa `st.pills`. */
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] {
                margin: 0 0 1rem !important;
            }
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] > [data-testid="stButtonGroup"] {
                display: flex !important;
                align-items: center !important;
                min-height: 2.85rem;
                padding: 0.3rem 0.45rem !important;
                background: #0a2745 !important;
                border: 1px solid #2c8396 !important;
                border-radius: 13px !important;
                box-shadow: 0 8px 20px rgba(0, 0, 0, 0.12);
            }
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radiogroup"] {
                display: flex !important;
                width: 100% !important;
                justify-content: center !important;
                gap: 0.2rem !important;
            }
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radio"] {
                flex: 0 1 auto !important;
                min-height: 2.25rem !important;
                padding: 0.48rem 0.95rem !important;
                border-color: transparent !important;
                border-radius: 8px !important;
                background: transparent !important;
                color: #c8eaf4 !important;
                -webkit-text-fill-color: #c8eaf4 !important;
                box-shadow: none !important;
            }
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radio"]:hover,
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radio"]:focus-visible,
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radio"][aria-checked="true"] {
                background: #176b8a !important;
                border-color: #6bbfd1 !important;
                color: #ffffff !important;
                -webkit-text-fill-color: #ffffff !important;
            }
            [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radio"] * {
                color: inherit !important;
                -webkit-text-fill-color: inherit !important;
            }
            [data-testid="stTabs"] [data-baseweb="tab-panel"] {
                padding-top: 0.55rem !important;
            }
            /* Las pestañas internas se leen como una segunda navegación, no
               como otro bloque grande de la página. */
            [data-testid="stTabs"] [data-baseweb="tab-panel"] [data-testid="stTabs"] {
                margin-top: 0.15rem !important;
                margin-bottom: 0.85rem !important;
            }
            [data-testid="stTabs"] [data-baseweb="tab-panel"] [data-testid="stTabs"] [role="tablist"] {
                padding: 0.25rem !important;
                background: rgba(13, 49, 83, 0.58) !important;
                border-color: rgba(107, 191, 209, 0.27) !important;
                box-shadow: none !important;
            }
            [data-testid="stTabs"] [data-baseweb="tab-panel"] [data-testid="stTabs"] [role="tab"] {
                min-height: 2.2rem !important;
                padding: 0.46rem 0.78rem !important;
                font-size: 0.82rem !important;
            }
            [data-testid="stWidgetLabel"] label,
            [data-testid="stWidgetLabel"] p {
                color: #c8eaf4 !important;
                font-size: 0.82rem !important;
                font-weight: 700 !important;
                letter-spacing: 0.01em;
            }
            [data-testid="stTextInput"],
            [data-testid="stSelectbox"],
            [data-testid="stFileUploader"] {
                filter: drop-shadow(0 4px 9px rgba(0, 0, 0, 0.08));
            }
            .stButton > button, .stDownloadButton > button {
                border-radius: 10px !important;
                box-shadow: 0 4px 11px rgba(0, 0, 0, 0.1);
                transition: transform 160ms ease, box-shadow 160ms ease,
                    background 160ms ease !important;
            }
            .stButton > button:hover:not(:disabled),
            .stDownloadButton > button:hover:not(:disabled) {
                transform: translateY(-1px);
                box-shadow: 0 7px 15px rgba(0, 0, 0, 0.16);
            }
            @media (max-width: 768px) {
                [data-testid="stMainBlockContainer"] {
                    width: 100% !important;
                    padding: 3.75rem 0.75rem 2rem !important;
                }
                .app-header {
                    align-items: flex-start;
                    grid-template-columns: 1fr;
                    margin-bottom: 0.8rem;
                    padding: 0.8rem 0.15rem 1rem;
                    border-radius: 0;
                }
                .app-title { font-size: 1.85rem !important; }
                .app-subtitle { font-size: 0.8rem !important; }
                [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] > [data-testid="stButtonGroup"] {
                    overflow-x: auto !important;
                    scrollbar-width: none;
                }
                [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] > [data-testid="stButtonGroup"]::-webkit-scrollbar {
                    display: none;
                }
                [data-testid="stElementContainer"][class*="st-key-analysis_model_mode"] [role="radiogroup"] {
                    width: max-content !important;
                    min-width: 100% !important;
                }
                [data-testid="stTabs"] [role="tablist"] {
                    flex-wrap: nowrap !important;
                    overflow-x: auto !important;
                    scrollbar-width: none;
                }
                [data-testid="stTabs"] [role="tablist"]::-webkit-scrollbar {
                    display: none;
                }
                [data-testid="stTabs"] [role="tab"] {
                    flex: 0 0 auto !important;
                    font-size: 0.8rem !important;
                    padding-inline: 0.75rem !important;
                }
            }
        </style>
        """,
        unsafe_allow_html=True,
    )
    analysis_mode = st.pills(
        "Sección principal",
        [WHITE_SPOT_LABEL, VACUOLIZATION_LABEL, HISTORY_LABEL],
        default=WHITE_SPOT_LABEL,
        key="analysis_model_mode",
        label_visibility="collapsed",
    )
    if analysis_mode != HISTORY_LABEL:
        st.markdown(
            """
            <header class="app-header">
                <div class="app-header-copy">
                    <div class="app-eyebrow">Análisis celular</div>
                    <h1 class="app-title">Monitoreo celular <span class="app-title-accent">en tiempo real</span></h1>
                    <p class="app-subtitle">
                        Detecta, revisa y guarda los resultados de tus muestras desde una sola vista.
                    </p>
                </div>
            </header>
            """,
            unsafe_allow_html=True,
        )

    # Parámetros internos: se ocultan para mantener la interfaz enfocada en el análisis.
    confidence = 0.35
    image_size = PHOTO_IMAGE_SIZE
    mask_opacity = 0.50
    low_limit, medium_limit, high_limit = 10.0, 30.0, 60.0

    model_path = default_model_path()
    if model_path is None and MODEL_BACKEND in {"openvino", "ov"}:
        st.session_state["openvino_fallback"] = True
        model_path = default_pt_model_path()
    if model_path is None or not (model_path.is_file() or is_openvino_model_dir(model_path)):
        if MODEL_BACKEND in {"openvino", "ov"}:
            st.error(
                "Se seleccionó OpenVINO, pero no se encontró el modelo convertido. "
                "Configura `CELL_OPENVINO_MODEL_DIR` o `CELL_OPENVINO_MODEL_URL`."
            )
        else:
            st.error("No se encontró el archivo del modelo `best.pt`.")
        st.stop()

    def get_app_model() -> YOLO:
        """Carga el modelo solo cuando el modo elegido realmente lo necesita."""
        nonlocal model_path
        try:
            loaded_model = get_session_model(model_path)
        except Exception as error:
            if MODEL_BACKEND in {"openvino", "ov"} and model_path.is_dir():
                st.session_state["openvino_fallback"] = True
                fallback_path = default_pt_model_path()
                try:
                    loaded_model = get_session_model(fallback_path)
                    st.warning(
                        "OpenVINO no pudo compilarse en este CPU; se está usando `.pt` como respaldo. "
                        "La app continúa funcionando, pero este equipo requiere otra optimización."
                    )
                    model_path = fallback_path
                except Exception as fallback_error:
                    st.error(
                        "No fue posible cargar el modelo OpenVINO ni el respaldo `.pt`: "
                        f"{fallback_error}"
                    )
                    st.stop()
            else:
                st.error(f"No fue posible cargar el modelo: {error}")
                st.stop()

        if loaded_model.task != "segment":
            st.error(
                f"El modelo cargado es de tipo `{loaded_model.task}`. "
                "Esta aplicación requiere un modelo de segmentación."
            )
            st.stop()
        return loaded_model

    def render_uploaded_video_mode() -> None:
        """Renderiza el análisis de un archivo de video dentro de su pestaña."""
        uploaded_video = st.file_uploader(
            "Carga un video para analizar",
            type=["mp4", "avi", "mov", "mkv"],
            key="uploaded_video_file",
        )
        if uploaded_video is None:
            st.info("Cuando tengas el video, súbelo aquí y presiona **Analizar y reproducir detección**.")
            return

        video_display = st.selectbox(
            "Visualización del video",
            ["Con máscaras", "Sin máscaras"],
            key="video_display",
        )
        show_video_masks = video_display == "Con máscaras"
        if not st.button("Analizar y reproducir detección", type="primary", use_container_width=True):
            return

        model = get_app_model()

        st.subheader("Detección en curso")
        st.caption(
            "El video original se muestra con las máscaras superpuestas mientras se analiza."
            if show_video_masks
            else "El video original se muestra mientras se analiza y se realiza el conteo."
        )
        preview_placeholder = st.empty()
        progress_bar = st.progress(0.0)
        status_text = st.empty()
        try:
            video_masks, counts, tracked_detections, video_measurements = process_video(
                model,
                uploaded_video,
                confidence,
                VIDEO_IMAGE_SIZE,
                mask_opacity,
                VIDEO_INFERENCE_STRIDE,
                show_video_masks,
                progress_bar,
                status_text,
                preview_placeholder,
            )
        except Exception as error:
            progress_bar.empty()
            status_text.empty()
            gc.collect()
            st.exception(error)
            return

        # Al terminar se retira la previsualización temporal: queda un solo video final.
        preview_placeholder.empty()
        status_text.success("Análisis terminado.")
        affected_percentage, grade, description = show_summary(
            counts,
            low_limit,
            medium_limit,
            high_limit,
            "Total de células únicas",
            sidebar=True,
            measurements=video_measurements,
        )
        if tracked_detections == 0:
            st.warning("No se obtuvieron identificadores de seguimiento; revisa la confianza, el video y el modelo.")

        st.subheader("Video original con máscaras" if show_video_masks else "Video original")
        st.video(video_masks, format="video/mp4")
        download_video, download_report = st.columns(2)
        with download_video:
            st.download_button(
                "Descargar video con máscaras" if show_video_masks else "Descargar video original",
                data=video_masks,
                file_name="celulas_con_mascaras.mp4" if show_video_masks else "video_original.mp4",
                mime="video/mp4",
                use_container_width=True,
            )
        with download_report:
            report = make_report_csv(
                counts,
                affected_percentage,
                grade,
                description,
                "Total de células únicas rastreadas",
                measurements=video_measurements,
            )
            st.download_button(
                "Descargar reporte CSV",
                data=report,
                file_name="reporte_celular.csv",
                mime="text/csv",
                use_container_width=True,
            )

        st.caption(
            "El conteo corresponde a IDs únicos seguidos a lo largo del video. "
            "La calificación debe validarse con tu criterio científico o clínico."
        )
        if video_measurements:
            st.caption(
                "Las medidas promedio se calcularon con las máscaras de los fotogramas analizados "
                f"({video_measurements['calibration_method']})."
            )

    if analysis_mode == WHITE_SPOT_LABEL:
        photo_tab, video_tab = st.tabs(["Foto", "Video"])
        with photo_tab:
            render_photo_mode(
                None,
                confidence,
                image_size,
                mask_opacity,
                low_limit,
                medium_limit,
                high_limit,
                model_loader=get_app_model,
            )
        with video_tab:
            live_camera_tab, uploaded_video_tab = st.tabs(["Cámara en vivo", "Subir video"])
            with live_camera_tab:
                # La carga/compilación de OpenVINO se inicia en segundo plano después
                # de pulsar «Iniciar detección». La vista WebRTC debe seguir dibujándose
                # para que capture fotogramas desde el primer clic.
                model = st.session_state.get("cell_model")

                def load_live_model_for_session() -> YOLO:
                    return load_live_model(model_path)

                render_live_camera(
                    model,
                    confidence,
                    mask_opacity,
                    model_loader=load_live_model_for_session,
                )
            with uploaded_video_tab:
                render_uploaded_video_mode()
    elif analysis_mode == VACUOLIZATION_LABEL:
        st.info("La sección de vacuolización está preparada para incorporar su modelo.")
    else:
        render_analysis_history()


if __name__ == "__main__":
    main()
